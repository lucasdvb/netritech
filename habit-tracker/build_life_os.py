#!/usr/bin/env python3
"""Build Lucas's Life OS habit tracker: a Google-Sheets-ready .xlsx plus an Apps Script finisher.

Styled after the Quiet Progress habit/task trackers (dark navy panels, colour-coded weeks,
check grids, analysis bars, area and line charts, daily donuts), rebuilt as a personal
operating system: priority levels, a Foundation Score, rolling 7/30-day consistency,
Minimum Viable Day and sick-day modes, adaptive steps, a 7-day weight average and
rule-based coaching.

    python3 build_life_os.py            # clean tracker  -> Lucas-Life-OS.xlsx + life_os_setup.gs
    python3 build_life_os.py --demo     # sample data and a fixed "today", for previews only
"""
import argparse
import datetime as dt
import os
import random

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- configuration
START_MONTH = (2026, 10)          # first month tab; the workbook covers 12 months
TRACK_START = dt.date(2026, 10, 7)
FIRST_MONDAY = dt.date(2026, 9, 28)  # Monday of the week that holds the first month's 1st
FONT = "Montserrat"
TODAY = "TODAY()"                 # replaced by a fixed date in --demo

# ---------------------------------------------------------------- palette (Quiet Progress dark theme)
BG = "141B2D"
PANEL = "1C2438"
PANEL2 = "232D45"
LINE = "2E3852"
TEXT = "E8ECF4"
MUTED = "8D97B0"
DIM = "5B6584"
VIOLET = "8B5CF6"
BLUE = "4F7CF7"
TEAL = "19C3C3"
PINK = "F0508A"
GREEN = "22C55E"
YELLOW = "F5C542"
ORANGE = "F59E42"
LIME = "A3E635"
FUCHSIA = "E879F9"
WEEK_COLORS = [VIOLET, BLUE, TEAL, PINK, GREEN]
DAY_COLORS = [VIOLET, BLUE, TEAL, GREEN, LIME, YELLOW, ORANGE]  # Mon..Sun
LEVEL_COLORS = {"L1": YELLOW, "L2": BLUE, "L3": MUTED}
LEVEL_TITLES = {
    "L1": "LEVEL 1 · NON-NEGOTIABLES · the daily foundation",
    "L2": "LEVEL 2 · HIGH VALUE · weekly targets",
    "L3": "LEVEL 3 · OPTIONAL · bonus, never guilt",
}
PILLARS = {
    "Body": PINK, "Health": TEAL, "Posture": BLUE, "Mind": VIOLET,
    "Spirit": YELLOW, "Work": ORANGE, "Relationships": FUCHSIA, "Life admin": LIME,
}

# ---------------------------------------------------------------- habits
# key, name, level, pillar, frequency, target, weekly target (0 = optional), what counts, MVD item
HABITS = [
    ("wake", "🌅 Morning reset", "L1", "Health", "Daily", "06:00", 7,
     "Up at 06:00 · 500–750 ml water · make bed · 5–15 min outdoor light · no socials for 30 min", False),
    ("pray", "🙏 Prayer + Scripture", "L1", "Spirit", "Daily", "5–10 min", 7,
     "Prayer, a short Scripture reading, one gratitude, one thing I need God's help with today", True),
    ("mob", "🧘 Mobility & posture", "L1", "Posture", "Daily", "10 min", 7,
     "Chin tucks · wall angels · thoracic extensions · external rotation · scap retractions · doorway stretch · cat-cow · hip flexors", True),
    ("move", "💪 Training / movement", "L1", "Body", "Daily", "per plan", 7,
     "The day's session from the plan, or a 20–30 min walk on recovery days", True),
    ("top3", "🎯 Top 3 priorities", "L1", "Work", "Workdays", "3 set", 5,
     "Before work: write the 3 things that make today a win", False),
    ("love", "❤️ Fiancée time", "L1", "Relationships", "Daily", "10–20 min", 7,
     "Intentional and phone-free. 'How are you really doing?'", True),
    ("read", "📖 Read / learn", "L1", "Mind", "Daily", "20 min", 7,
     "20 min reading or deliberate learning (AI, marketing, design, sales, leadership, faith)", False),
    ("shut", "🌙 Evening shutdown", "L1", "Work", "Daily", "20:00 / 21:00", 7,
     "Work closed by 20:00 · clothes + kit ready · tomorrow reviewed · last 30 min screen-free", False),
    ("sleep", "😴 Lights out 22:00", "L1", "Health", "Daily", "22:00", 7,
     "Lights out around 22:00 (7.5–8.5 h before a 06:00 wake)", True),
    ("str", "🏋️ Strength session", "L2", "Body", "4× / wk", "45–60 min", 4,
     "Mon/Thu upper + posture · Tue/Fri lower + calves + core", False),
    ("prog", "📈 Progression achieved", "L2", "Body", "Each session", "yes / no", 4,
     "+1 rep, +1 set, better form, slower lowering, longer hold, harder variation, less rest, more range or heavier", False),
    ("core", "🧱 Core", "L2", "Body", "3× / wk", "10–15 min", 3,
     "Plank · side plank · dead bug · hollow hold · reverse crunch · leg-raise progression", False),
    ("calf", "🦵 Calves", "L2", "Body", "3× / wk", "10 min", 3,
     "Standing → single-leg → slow eccentric → paused calf raises", False),
    ("cardio", "🚶 Cardio", "L2", "Body", "2× / wk", "20–30 min", 2,
     "Brisk or incline walk, cycling, easy jog, low-impact cardio", False),
    ("deep", "🧠 Deep work 2+ blocks", "L2", "Work", "Workdays", "2–3 blocks", 5,
     "60–90 min focus blocks with notifications off", False),
    ("brk", "👀 Posture & eye breaks", "L2", "Posture", "Workdays", "6–8 / day", 5,
     "Move 1–2 min every 45–60 min · 20-20-20 for the eyes · don't rub eyes", False),
    ("home", "🧹 Home reset", "L2", "Life admin", "5× / wk", "5 min", 5,
     "5-minute tidy plus dishes/kitchen reset", False),
    ("son", "👦 Father–son time", "L2", "Relationships", "Weekly", "1+", 1,
     "Conversation, shared activity, guidance or positive reinforcement", False),
    ("couple", "🥂 Couple time / date", "L2", "Relationships", "Weekly", "1×", 1,
     "Dinner, walk, coffee, movie, day trip, private time", False),
    ("family", "📞 Family check-in", "L2", "Relationships", "Weekly", "1×", 1,
     "A call or visit with one meaningful conversation", False),
    ("church", "⛪ Church / worship", "L2", "Spirit", "Weekly", "1×", 1,
     "Sunday worship or equivalent", False),
    ("prep", "🍱 Meal prep", "L2", "Life admin", "Weekly", "30–60 min", 1,
     "Sunday: protein, carb base, veg, easy snacks", False),
    ("med", "🕊️ Meditation / journal", "L3", "Mind", "Optional", "5–10 min", 0,
     "Breath awareness, contemplative prayer or the three journal questions", False),
    ("xcardio", "➕ Extra walk / cardio", "L3", "Body", "Optional", "bonus", 0, "Bonus only", False),
    ("xmob", "➕ Extra mobility", "L3", "Posture", "Optional", "bonus", 0, "Bonus only", False),
    ("xread", "➕ Extra reading", "L3", "Mind", "Optional", "bonus", 0, "Bonus only", False),
]
H = {h[0]: h for h in HABITS}
HKEYS = [h[0] for h in HABITS]
L1KEYS = [h[0] for h in HABITS if h[2] == "L1"]

# key, label, unit, number format, target text
METRICS = [
    ("weight", "⚖️ Weight", "kg", "0.0", "AM, 7-day avg"),
    ("steps", "👣 Steps", "steps", "0", "adaptive"),
    ("protein", "🍗 Protein", "g", "0", "150 g"),
    ("water", "💧 Water", "L", "0.0", "2.5–3 L"),
    ("kcal", "🔥 Calories", "kcal", "0", "1.8–2k"),
    ("sleeph", "🛌 Sleep", "hours", "0.0", "7.5–8.5 h"),
    ("energy", "⚡ Energy", "1–10", "0", "1–10"),
    ("stress", "🌡️ Stress", "1–10", "0", "1–10"),
    ("mood", "🙂 Mood", "1–10", "0", "1–10"),
]
MKEYS = [m[0] for m in METRICS]

# Auto-scored items that join the habit library (derived from the numbers you log)
AUTO = [
    ("c_steps", "👣 Steps target hit", "L1", "Body", "Daily", "7k → 9k", 7, "Adaptive: 7,000 in week 1, 8,000 in week 2, then 9,000"),
    ("c_prot", "🍗 Protein ≥ 135 g", "L1", "Health", "Daily", "150 g", 7, "Counts as hit at 90% of the 150 g target"),
    ("c_water", "💧 Water ≥ 2.5 L", "L1", "Health", "Daily", "2.5–3 L", 7, "500 ml × 5–6 bottles"),
    ("c_sleep", "🛌 Sleep ≥ 7 h", "L1", "Health", "Daily", "7.5–8.5 h", 7, "Uses sleep hours when logged, otherwise the lights-out tick"),
]

# ---------------------------------------------------------------- data engine layout
DATA_FIRST = 4
DATA_LAST = DATA_FIRST + 12 * 31 - 1
DATA_KEYS = (["date", "month", "day", "wd", "mode"] + HKEYS + MKEYS +
             ["active", "logged", "wk", "stepT", "workday", "strday",
              "c_sleep", "c_move", "c_prot", "c_water", "c_steps", "c_pray", "c_work", "c_love", "c_read",
              "norm", "mvd", "score", "r7", "r30", "w7", "l1done", "l1app",
              "score_na", "w_na", "w7_na", "en_na", "st_na", "mo_na"])
assert len(set(DATA_KEYS)) == len(DATA_KEYS), "duplicate data keys"
DC = {k: CL(i + 1) for i, k in enumerate(DATA_KEYS)}


def DR(key):
    c = DC[key]
    return f"Data!${c}${DATA_FIRST}:${c}${DATA_LAST}"


def drow(k, j):
    return DATA_FIRST + k * 31 + (j - 1)


def months():
    y, m = START_MONTH
    out = []
    for k in range(12):
        mm = (m - 1 + k) % 12 + 1
        yy = y + (m - 1 + k) // 12
        out.append((k, yy, mm, dt.date(yy, mm, 1).strftime("%b %Y")))
    return out


MONTHS = months()

# ---------------------------------------------------------------- style helpers
THIN_BG = Side(style="thin", color=BG)
TILE = Border(left=THIN_BG, right=THIN_BG, top=THIN_BG, bottom=THIN_BG)
CHECKBOXES = {}   # sheet -> [A1 ranges] for the Apps Script
CHART_STYLES = {}  # sheet -> list of chart style dicts in creation order


def mix(a, b, t):
    """Blend colour a into b by t (0 = b, 1 = a)."""
    a = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x * t + y * (1 - t)):02X}" for x, y in zip(a, b))


def PF(c):
    return PatternFill("solid", start_color=c, end_color=c)


def F(size=9, color=TEXT, bold=False, italic=False):
    return Font(name=FONT, size=size, color=color, bold=bold, italic=italic)


def cells(ws, rng):
    for row in ws[rng]:
        for c in row:
            yield c


def paint(ws, rng, bg):
    for c in cells(ws, rng):
        c.fill = PF(bg)


def put(ws, ref, value=None, *, size=9, color=TEXT, bold=False, italic=False, bg=None,
        h="left", v="center", wrap=False, fmt=None, indent=0, border=None):
    """Write and style a cell, or a range (merged) when ref contains ':'."""
    rng = ref if ":" in ref else None
    if rng:
        for c in cells(ws, rng):
            if bg:
                c.fill = PF(bg)
            if border:
                c.border = border
        ws.merge_cells(rng)
        ref = rng.split(":")[0]
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = F(size, color, bold, italic)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)
    if bg:
        c.fill = PF(bg)
    if fmt and not (fmt == "@" and isinstance(c.value, str) and c.value.startswith("=")):
        c.number_format = fmt
    if border:
        c.border = border
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def heights(ws, rows, h):
    for r in rows:
        ws.row_dimensions[r].height = h


def cf(ws, rng, formula, fill=None, color=None, bold=None, stop=False):
    font = Font(color=color, bold=bold) if (color or bold) else None
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[formula], fill=PF(fill) if fill else None, font=font, stopIfTrue=stop))


def band_cf(ws, rng, ref, mode="font"):
    """Colour by the status bands (>=85 green, >=70 teal, >=50 yellow, else pink)."""
    for lo, col in ((0.85, GREEN), (0.70, TEAL), (0.50, YELLOW)):
        f = f'AND(ISNUMBER({ref}),{ref}>={lo})'
        cf(ws, rng, f, color=col, stop=True) if mode == "font" else cf(ws, rng, f, fill=mix(col, PANEL, .35), stop=True)
    f = f'AND(ISNUMBER({ref}),{ref}<0.5)'
    cf(ws, rng, f, color=PINK, stop=True) if mode == "font" else cf(ws, rng, f, fill=mix(PINK, PANEL, .35), stop=True)


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def bar(x, n=10):
    return (f'IF(ISNUMBER({x}),REPT("█",ROUND(MAX(0,MIN(1,{x}))*{n},0))&'
            f'REPT("░",{n}-ROUND(MAX(0,MIN(1,{x}))*{n},0)),"")')


def status(x):
    return (f'IF(NOT(ISNUMBER({x})),"—",IF({x}>=Band1,"GREEN",IF({x}>=Band2,"GOOD",'
            f'IF({x}>=Band3,"NEEDS ATTENTION","RESET"))))')


def status_cf(ws, rng, ref):
    for word, col in (("GREEN", GREEN), ("GOOD", TEAL), ("NEEDS ATTENTION", YELLOW), ("RESET", PINK)):
        cf(ws, rng, f'{ref}="{word}"', color=col, stop=True)


def page(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def rich(color=MUTED, size=800, bold=False):
    cp = CharacterProperties(sz=size, b=bold, solidFill=color, latin=DFont(typeface=FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def style_chart(ch, *, ymin=None, ymax=None, yfmt=None, major=None, legend=True, bg=PANEL):
    ch.title = None
    ch.graphical_properties = GraphicalProperties(solidFill=bg, ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(solidFill=bg, ln=LineProperties(noFill=True))
    ch.display_blanks = "gap"
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.txPr = rich(MUTED, 700)
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=LINE)))
    ch.x_axis.majorGridlines = None
    if ymin is not None:
        ch.y_axis.scaling.min = ymin
    if ymax is not None:
        ch.y_axis.scaling.max = ymax
    if major:
        ch.y_axis.majorUnit = major
    if yfmt:
        ch.y_axis.numFmt = NumFmt(formatCode=yfmt, sourceLinked=False)
    if legend:
        ch.legend.position = "b"
        ch.legend.txPr = rich(MUTED, 800)
    else:
        ch.legend = None


def line_series(ch, color, width=22000, smooth=True, dash=None):
    s = ch.series[-1]
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = width
    if dash:
        s.graphicalProperties.line.dashStyle = dash
    s.smooth = smooth
    s.marker.symbol = "none"


def size_chart(ch, ws, c1, c2, r1, r2):
    """Size a chart to cover columns c1..c2 and rows r1..r2 (inclusive)."""
    wpx = sum((ws.column_dimensions[CL(c)].width or 8.43) * 7 for c in range(c1, c2 + 1))
    hpx = sum((ws.row_dimensions[r].height or 15) * 4 / 3 for r in range(r1, r2 + 1))
    ch.width = wpx / 37.8
    ch.height = hpx / 37.8
    ch.anchor = f"{CL(c1)}{r1}"
    ch._life_os_anchor = (r1, c1)


def chart_style(sheet, ch, kind, colors, **extra):
    row, col = ch._life_os_anchor
    CHART_STYLES.setdefault(sheet, []).append(
        {"kind": kind, "row": row, "col": col, "colors": ["#" + c for c in colors], **extra})


def checkbox(sheet, rng):
    CHECKBOXES.setdefault(sheet, []).append(rng)


# ======================================================================== SETTINGS
SETTINGS = [
    ("PROFILE", None),
    ("Name", "OwnerName", "Lucas", "@", ""),
    ("Age", "Age", 35, "0", ""),
    ("Height (cm)", "HeightCm", 180, "0", ""),
    ("Starting weight (kg)", "StartWeight", 76.0, "0.0", "Used for goal estimates only"),
    ("Estimated body fat", "StartBF", 0.25, "0%", "Rough estimate"),
    ("Goal body fat", "GoalBF", 0.15, "0%", "Direction of travel, not a deadline"),
    ("Estimated lean mass (kg)", "LeanMass", "=StartWeight*(1-StartBF)", "0.0", "Calculated"),
    ("Goal weight if muscle is kept (kg)", "GoalWeight", "=LeanMass/(1-GoalBF)", "0.0", "Calculated estimate"),
    ("SCHEDULE", None),
    ("Tracking start date", "TrackStart", TRACK_START, "d mmm yyyy", "Days before this are not scored"),
    ("Wake time", "WakeTime", "06:00", "@", ""),
    ("Training time", "TrainTime", "06:30", "@", ""),
    ("Work hours", "WorkHours", "10:00–20:00", "@", ""),
    ("Lights out", "LightsOut", "22:00", "@", ""),
    ("Work days (1 = Mon … 7 = Sun)", "WorkDays", "12345", "@", "Top 3 only counts on these days"),
    ("Strength days (1 = Mon … 7 = Sun)", "StrDays", "1245", "@", "Mon, Tue, Thu, Fri"),
    ("NUTRITION", None),
    ("Protein target (g)", "ProtTarget", 150, "0", "Priority #1 nutrition habit"),
    ("Protein counted as hit (g)", "ProtHit", 135, "0", "90% of target"),
    ("Protein alert before dinner (g)", "ProtDinner", 120, "0", "Below this → high-protein dinner"),
    ("Protein on a Minimum Viable Day (g)", "ProtMVD", 120, "0", "As close as possible"),
    ("Calories, low end (kcal)", "KcalMin", 1800, "0", "Start here, adjust on the 2–3 week trend"),
    ("Calories, high end (kcal)", "KcalMax", 2000, "0", "Never auto-drop to 1,500"),
    ("Fat (g/day)", "FatRange", "50–65", "@", "Portion awareness, not fear"),
    ("Water counted as hit (L)", "WaterHit", 2.5, "0.0", "Target 2.5–3 L, more in heat"),
    ("Water on a Minimum Viable Day (L)", "WaterMVD", 2.0, "0.0", ""),
    ("SLEEP", None),
    ("Sleep target (h)", "SleepTarget", 7.5, "0.0", "7.5–8.5 h"),
    ("Sleep counted as hit (h)", "SleepMin", 7.0, "0.0", ""),
    ("Low-sleep alert (h)", "SleepLow", 6.0, "0.0", "Below this → recovery day"),
    ("STEPS (adaptive)", None),
    ("Week 1 target", "Steps1", 7000, "#,##0", ""),
    ("Week 2 target", "Steps2", 8000, "#,##0", ""),
    ("Week 3 onward", "Steps3", 9000, "#,##0", "Aim for a 9–10k average"),
    ("Evening alert (steps by 18:00)", "StepsAlert", 5000, "#,##0", "Below this → 20–30 min walk"),
    ("BODY", None),
    ("Weekly loss, low end (kg)", "LossMin", 0.4, "0.0", ""),
    ("Weekly loss, high end (kg)", "LossMax", 0.8, "0.0", "Not a fixed 1 kg/week"),
    ("Too fast (% bodyweight / week)", "LossMaxPct", 0.01, "0.0%", "Repeatedly above → eat more"),
    ("SCORING", None),
    ("Green from", "Band1", 0.85, "0%", "≥ 85%"),
    ("Good from", "Band2", 0.70, "0%", "70–84%"),
    ("Needs attention from", "Band3", 0.50, "0%", "50–69%, below is Reset"),
    ("Recovery score alert", "RecoveryLow", 60, "0", "Below this → lighter training"),
]

IDEAL_WEEK = [
    ("Monday", "Upper body + posture", "Work 10–20 · family 20:00"),
    ("Tuesday", "Lower body + calves + core", "Work 10–20 · family 20:00"),
    ("Wednesday", "Walk + mobility (recovery)", "Work 10–20 · family 20:00"),
    ("Thursday", "Upper body + posture", "Work 10–20 · family 20:00"),
    ("Friday", "Lower body + calves + core", "Work 10–20 · family 20:00"),
    ("Saturday", "Optional cardio / walk / mobility", "Family · fiancée · life"),
    ("Sunday", "Recovery · church", "Meal prep · weekly review · plan"),
]


def build_settings(wb, ws):
    page(ws)
    ws.sheet_properties.tabColor = DIM
    paint(ws, "A1:L60", BG)
    widths(ws, {"A": 2, "B": 36, "C": 16, "D": 40, "E": 3, "F": 6, "G": 14, "H": 34, "I": 32, "J": 3})
    put(ws, "B2", "SETTINGS", size=22, bold=True)
    put(ws, "B3", "Targets and thresholds behind every score. Change a value here and the whole system follows.",
        size=9, color=MUTED)
    r = 5
    for row in SETTINGS:
        if row[1] is None:
            put(ws, f"B{r}:D{r}", row[0], size=8, bold=True, color=YELLOW, bg=PANEL2, indent=1)
            r += 1
            continue
        label, nm, val, fmt, note = row
        put(ws, f"B{r}", label, bg=PANEL, indent=1)
        c = put(ws, f"C{r}", val, bg=PANEL2, h="center", bold=True, color=YELLOW if not str(val).startswith("=") else TEXT,
                fmt=fmt)
        if fmt == "@":
            c.number_format = "@"
        put(ws, f"D{r}", note, size=8, color=MUTED, bg=PANEL, indent=1)
        name(wb, nm, f"Settings!$C${r}")
        r += 1
    put(ws, "B" + str(r + 1), "Yellow values are yours to edit. White values are calculated.", size=8, color=MUTED,
        italic=True)

    put(ws, "F5:I5", "IDEAL WEEK", size=8, bold=True, color=YELLOW, bg=PANEL2, indent=1)
    for i, hd in enumerate(["#", "Day", "06:30 training", "Rest of the day"]):
        put(ws, f"{'FGHI'[i]}6", hd, size=8, bold=True, color=MUTED, bg=PANEL, h="center" if i == 0 else "left",
            indent=0 if i == 0 else 1)
    for i, (day, train, rest) in enumerate(IDEAL_WEEK):
        rr = 7 + i
        put(ws, f"F{rr}", i + 1, bg=PANEL, h="center", color=DAY_COLORS[i], bold=True)
        put(ws, f"G{rr}", day, bg=PANEL, color=DAY_COLORS[i], bold=True, indent=1)
        put(ws, f"H{rr}", train, bg=PANEL2, color=YELLOW, indent=1)
        put(ws, f"I{rr}", rest, bg=PANEL2, color=YELLOW, indent=1)
    name(wb, "PlanTraining", "Settings!$H$7:$H$13")
    name(wb, "PlanRest", "Settings!$I$7:$I$13")
    put(ws, "F15:I15", "DAILY TIMELINE (weekdays)", size=8, bold=True, color=YELLOW, bg=PANEL2, indent=1)
    timeline = [("06:00", "Wake · water · bed · light · 2–5 min quiet/prayer"),
                ("06:05", "Morning reset: water, outdoor light, prayer + Scripture"),
                ("06:15", "Mobility & posture, 10 min"),
                ("06:30", "Training per the ideal week"),
                ("07:30", "Shower · breakfast (30–40 g protein)"),
                ("09:30", "Top 3 priorities · optional journal"),
                ("10:00", "Work: 2–3 deep-work blocks, breaks every 45–60 min"),
                ("20:00", "Work shutdown → fiancée, son, family"),
                ("21:00", "Shutdown routine · screens down"),
                ("21:30", "Read / quiet time"),
                ("22:00", "Lights out")]
    for i, (t, what) in enumerate(timeline):
        rr = 16 + i
        put(ws, f"F{rr}:G{rr}", t, bg=PANEL, bold=True, color=TEAL, h="center")
        put(ws, f"H{rr}:I{rr}", what, bg=PANEL, indent=1)
    heights(ws, range(1, 60), 18)
    ws.row_dimensions[2].height = 34


# ======================================================================== MONTH SHEETS
MC_FIRST = 5                       # column E = day 1
MC_LAST = MC_FIRST + 30            # column AI = day 31
AK, ALc, AM, AN, AP = "AK", "AL", "AM", "AN", "AP"
R_MODE = 8
R_SCORE = 39
R_L1 = 40
R_METRIC_HDR = 56
R_METRIC0 = 59
R_W7 = R_METRIC0 + len(METRICS)       # 68
R_STEPT = R_W7 + 1                     # 69
WEEK_BLOCKS = [(1, 7), (8, 14), (15, 21), (22, 28), (29, 31)]


def month_rows():
    rows, bands, r = {}, {}, 9
    for lvl in ("L1", "L2", "L3"):
        bands[lvl] = r
        r += 1
        for h in HABITS:
            if h[2] == lvl:
                rows[h[0]] = r
                r += 1
    return rows, bands


HROW, BANDROW = month_rows()
MROW = {m[0]: R_METRIC0 + i for i, m in enumerate(METRICS)}


def dcol(j):
    return CL(MC_FIRST + j - 1)


def habit_ranges():
    """Contiguous habit-grid row ranges per level."""
    out = []
    for lvl in ("L1", "L2", "L3"):
        rs = [HROW[h[0]] for h in HABITS if h[2] == lvl]
        out.append((min(rs), max(rs)))
    return out


def build_month(wb, ws, k, y, m):
    sn = ws.title
    page(ws)
    ws.sheet_properties.tabColor = WEEK_COLORS[k % 5]
    paint(ws, f"A1:AO{R_STEPT + 18}", BG)
    widths(ws, {"A": 2, "B": 27, "C": 11, "D": 9, "AJ": 2, "AK": 7, "AL": 7, "AM": 15, "AN": 7, "AO": 2, "AP": 3})
    for j in range(1, 32):
        ws.column_dimensions[dcol(j)].width = 3.7
    heights(ws, range(1, R_STEPT + 20), 16)
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 8

    # hidden helpers (same-sheet so Google conditional formatting can read them)
    helpers = {1: f"=DATE({y},{m},1)", 2: "=ProtHit", 3: "=WaterHit", 4: "=SleepMin", 5: "=SleepLow",
               6: f"={TODAY}", 7: "=KcalMin", 8: "=KcalMax",
               9: f'=COUNTIFS({DR("month")},{k + 1},{DR("active")},1,{DR("mode")},"<>S")',
               10: f'=COUNTIFS({DR("month")},{k + 1},{DR("active")},1,{DR("mode")},"<>S",{DR("workday")},1)',
               11: f'=_xlfn.MAXIFS({DR("date")},{DR("month")},{k + 1},{DR("weight")},">0")',
               12: f'=_xlfn.MINIFS({DR("date")},{DR("month")},{k + 1},{DR("weight")},">0")',
               13: f"=MIN(AP6,EOMONTH(AP1,0))"}
    for r, f in helpers.items():
        c = ws[f"AP{r}"]
        c.value = f
        c.font = F(6, BG)
        c.fill = PF(BG)

    # ---- title + KPI cards
    put(ws, "B2:D2", "=UPPER(TEXT($AP$1,\"mmmm\"))", size=22, bold=True)
    put(ws, "B3:D3", '="— Life OS · Habit Tracker · "&YEAR($AP$1)&" —"', size=8, color=MUTED, v="top")
    hr = habit_ranges()
    grid_sum = "+".join(f"COUNTIF(E{a}:AI{b},TRUE)+COUNTIF(E{a}:AI{b},\"x\")" for a, b in hr)
    month_r7 = f'IFERROR(INDEX({DR("r7")},MATCH($AP$13,{DR("date")},0)),"")'
    month_r30 = f'IFERROR(INDEX({DR("r30")},MATCH($AP$13,{DR("date")},0)),"")'
    cards = [
        ("E", "J", "HABITS", f"=COUNTA(B{hr[0][0]}:B{hr[0][1]},B{hr[1][0]}:B{hr[1][1]},B{hr[2][0]}:B{hr[2][1]})", TEXT),
        ("L", "Q", "CHECK-INS", f"={grid_sum}", TEXT),
        ("S", "AA", "FOUNDATION · MONTH", f'=IF(ISNUMBER(AK{R_SCORE}),TEXT(AK{R_SCORE},"0%")&"   "&{bar(f"AK{R_SCORE}", 12)},"—")', YELLOW),
        ("AC", "AI", "ROLLING 7-DAY", f'=IF(ISNUMBER({month_r7}),TEXT({month_r7},"0%")&"  ·  "&{status(month_r7)},"—")', TEAL),
        ("AK", "AN", "ROLLING 30-DAY", f'=IF(ISNUMBER({month_r30}),TEXT({month_r30},"0%")&"  ·  "&{status(month_r30)},"—")', TEAL),
    ]
    for c1, c2, label, f, col in cards:
        put(ws, f"{c1}2:{c2}2", label, size=7, bold=True, color=MUTED, bg=PANEL, h="center", v="bottom")
        put(ws, f"{c1}3:{c2}3", f, size=11, bold=True, color=col, bg=PANEL, h="center")

    # ---- week headers, weekdays, dates
    def week_header(row, title):
        put(ws, f"B{row}:D{row}", title, size=10, bold=True, color=TEXT, bg=PANEL, indent=1)
        for w, (a, b) in enumerate(WEEK_BLOCKS):
            put(ws, f"{dcol(a)}{row}:{dcol(b)}{row}", f"Week {w + 1}", size=8, bold=True, color="FFFFFF",
                bg=WEEK_COLORS[w], h="center")
        put(ws, f"AK{row}:AN{row}", "ANALYSIS", size=9, bold=True, color=TEXT, bg=PANEL, h="center")

    def day_rows(r_dow, r_date, src=None):
        for j in range(1, 32):
            c = dcol(j)
            w = next(i for i, (a, b) in enumerate(WEEK_BLOCKS) if a <= j <= b)
            tint = mix(WEEK_COLORS[w], PANEL, .28)
            if src is None:
                ws[f"{c}{r_date}"] = f'=IF({j}<=DAY(EOMONTH($AP$1,0)),$AP$1+{j - 1},"")'
                ws[f"{c}{r_dow}"] = f'=IF({c}{r_date}="","",LEFT(TEXT({c}{r_date},"ddd"),2))'
            else:
                ws[f"{c}{r_date}"] = f"={c}{src[1]}"
                ws[f"{c}{r_dow}"] = f"={c}{src[0]}"
            put(ws, f"{c}{r_dow}", size=6, color=mix(WEEK_COLORS[w], TEXT, .35), bg=tint, h="center", border=TILE)
            put(ws, f"{c}{r_date}", size=7, bold=True, color=TEXT, bg=tint, h="center", fmt="d", border=TILE)

    week_header(5, "MY HABITS")
    day_rows(6, 7)
    put(ws, "B6:D6", "Tick what you did. Missed is data, not failure.", size=7, color=MUTED, bg=PANEL, indent=1)
    put(ws, "B7", "Habit", size=7, bold=True, color=MUTED, bg=PANEL, indent=1)
    put(ws, "C7", "Pillar", size=7, bold=True, color=MUTED, bg=PANEL)
    put(ws, "D7", "Target", size=7, bold=True, color=MUTED, bg=PANEL)
    for col, t in zip(("AK", "AL", "AM", "AN"), ("Done", "Goal", "Progress", "%")):
        put(ws, f"{col}6", t, size=7, bold=True, color=MUTED, bg=PANEL, h="center")
    put(ws, "AK7:AN7", "goal = weekly target × days tracked", size=6, color=DIM, bg=PANEL, h="center")

    # ---- day mode row
    put(ws, f"B{R_MODE}", "Day mode", size=8, bold=True, color=TEXT, bg=PANEL2, indent=1)
    put(ws, f"C{R_MODE}:D{R_MODE}", "M = minimum · S = sick · R = rest", size=6, color=MUTED, bg=PANEL2)
    for j in range(1, 32):
        put(ws, f"{dcol(j)}{R_MODE}", size=7, bold=True, color=BG, bg=PANEL2, h="center", border=TILE)
    put(ws, f"AK{R_MODE}:AN{R_MODE}",
        f'=COUNTIF(E{R_MODE}:AI{R_MODE},"M")&" min · "&COUNTIF(E{R_MODE}:AI{R_MODE},"S")&" sick · "&COUNTIF(E{R_MODE}:AI{R_MODE},"R")&" rest"',
        size=7, color=MUTED, bg=PANEL2, h="center")
    dv = DataValidation(type="list", formula1='"M,S,R"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Day mode", error="Use M (minimum viable day), S (sick) or R (rest). Leave blank for a normal day.")
    ws.add_data_validation(dv)
    dv.add(f"E{R_MODE}:AI{R_MODE}")

    # ---- level bands + habit rows
    for lvl, r in BANDROW.items():
        col = LEVEL_COLORS[lvl]
        put(ws, f"B{r}:D{r}", LEVEL_TITLES[lvl].split(" · ", 1)[0] + " · " + LEVEL_TITLES[lvl].split(" · ")[1],
            size=7, bold=True, color=col, bg=mix(col, BG, .14), indent=1)
        paint(ws, f"E{r}:AI{r}", mix(col, BG, .14))
        put(ws, f"AK{r}:AN{r}", LEVEL_TITLES[lvl].split(" · ")[2], size=6, italic=True, color=col,
            bg=mix(col, BG, .14), h="center")
        ws.row_dimensions[r].height = 13
    for key, nm, lvl, pillar, freq, target, weekly, what, mvd in HABITS:
        r = HROW[key]
        put(ws, f"B{r}", nm + ("  ◆" if mvd else ""), size=8, bg=PANEL, indent=1)
        put(ws, f"C{r}", pillar.upper(), size=6, bold=True, color=PILLARS[pillar], bg=PANEL)
        put(ws, f"D{r}", freq if weekly in (0, 7) else f"{weekly}× / wk", size=6, color=MUTED, bg=PANEL)
        for j in range(1, 32):
            w = next(i for i, (a, b) in enumerate(WEEK_BLOCKS) if a <= j <= b)
            put(ws, f"{dcol(j)}{r}", size=9, bold=True, color=WEEK_COLORS[w], bg=PANEL2, h="center", border=TILE)
        done = f"=COUNTIF(E{r}:AI{r},TRUE)+COUNTIF(E{r}:AI{r},\"x\")"
        put(ws, f"AK{r}", done, size=8, bold=True, bg=PANEL, h="center")
        if weekly == 0:
            put(ws, f"AL{r}", "—", size=8, color=DIM, bg=PANEL, h="center")
            put(ws, f"AM{r}", f'=IF(AK{r}=0,"",AK{r}&" × bonus")', size=7, color=MUTED, bg=PANEL, h="center")
            put(ws, f"AN{r}", None, bg=PANEL)
        else:
            goal = "$AP$10" if key == "top3" else f"ROUND($AP$9*{weekly}/7,0)"
            put(ws, f"AL{r}", f"={goal}", size=8, color=MUTED, bg=PANEL, h="center")
            put(ws, f"AN{r}", f'=IF(N(AL{r})=0,"",MIN(1,AK{r}/AL{r}))', size=8, bold=True, bg=PANEL, h="center",
                fmt="0%")
            put(ws, f"AM{r}", f"={bar(f'AN{r}', 10)}", size=8, color=TEAL, bg=PANEL)
    for a, b in hr:
        checkbox(sn, f"E{a}:AI{b}")

    # ---- score rows
    put(ws, f"B{R_SCORE}", "Foundation score", size=8, bold=True, color=YELLOW, bg=PANEL, indent=1)
    put(ws, f"C{R_SCORE}:D{R_SCORE}", "auto · 9 pillars of the day", size=6, color=MUTED, bg=PANEL)
    put(ws, f"B{R_L1}", "Non-negotiables done", size=8, color=TEXT, bg=PANEL, indent=1)
    put(ws, f"C{R_L1}:D{R_L1}", "of 9 (8 on non-work days)", size=6, color=MUTED, bg=PANEL)
    for j in range(1, 32):
        c, r = dcol(j), drow(k, j)
        put(ws, f"{c}{R_SCORE}", f"=Data!{DC['score']}{r}", size=6, bold=True, color=TEXT, bg=PANEL, h="center",
            fmt="0%", border=TILE)
        put(ws, f"{c}{R_L1}", f'=IF(Data!{DC["score"]}{r}="","",Data!{DC["l1done"]}{r})', size=7, color=MUTED,
            bg=PANEL, h="center", border=TILE)
    put(ws, f"AK{R_SCORE}", f'=IFERROR(AVERAGE(E{R_SCORE}:AI{R_SCORE}),"")', size=8, bold=True, bg=PANEL, h="center",
        fmt="0%")
    put(ws, f"AL{R_SCORE}", "avg", size=7, color=MUTED, bg=PANEL, h="center")
    put(ws, f"AN{R_SCORE}", f"=AK{R_SCORE}", size=8, bold=True, bg=PANEL, h="center", fmt="0%")
    put(ws, f"AM{R_SCORE}", f"={bar(f'AN{R_SCORE}', 10)}", size=8, color=YELLOW, bg=PANEL)
    paint(ws, f"AK{R_L1}:AN{R_L1}", PANEL)

    # ---- foundation chart block
    put(ws, "B42:D42", "DAILY FOUNDATION SCORE", size=9, bold=True, color=YELLOW, bg=PANEL, indent=1)
    put(ws, "B43:D54",
        "One number for the day, built from 9 things:\n\nSleep · Movement · Protein · Water · Steps · "
        "Prayer · Top 3 (work days) · Fiancée · Read/learn\n\nM days are scored on the Minimum Viable Day.\n"
        "S days are paused — no score, no guilt.\n\n◆ = part of the Minimum Viable Day",
        size=7, color=MUTED, bg=PANEL, wrap=True, v="top", indent=1)
    paint(ws, "E42:AI54", PANEL)
    put(ws, "AK42:AN42", "WEEKLY AVERAGE", size=8, bold=True, color=TEXT, bg=PANEL, h="center")
    paint(ws, "AK43:AN54", PANEL)
    for w, (a, b) in enumerate(WEEK_BLOCKS):
        r = 44 + w * 2
        put(ws, f"AK{r}:AL{r}", f"Week {w + 1}", size=8, bold=True, color=WEEK_COLORS[w], bg=PANEL, h="center")
        put(ws, f"AN{r}", f'=IFERROR(AVERAGE({dcol(a)}{R_SCORE}:{dcol(b)}{R_SCORE}),"")', size=8, bold=True,
            bg=PANEL, h="center", fmt="0%")
        put(ws, f"AM{r}", f"={bar(f'AN{r}', 10)}", size=8, color=WEEK_COLORS[w], bg=PANEL)
    ch = AreaChart()
    ch.grouping = "standard"
    a, b = drow(k, 1), drow(k, 31)
    data_ws = wb["Data"]
    ch.add_data(Reference(data_ws, min_col=DATA_KEYS.index("score_na") + 1, min_row=a, max_row=b), titles_from_data=False)
    ch.set_categories(Reference(data_ws, min_col=DATA_KEYS.index("day") + 1, min_row=a, max_row=b))
    s = ch.series[0]
    s.graphicalProperties.solidFill = mix(TEAL, PANEL, .45)
    s.graphicalProperties.line.solidFill = TEAL
    s.graphicalProperties.line.width = 22000
    style_chart(ch, ymin=0, ymax=1, yfmt="0%", major=0.25, legend=False)
    size_chart(ch, ws, MC_FIRST, MC_LAST, 43, 54)
    ws.add_chart(ch)
    chart_style(sn, ch, "area", [TEAL], vmin=0, vmax=1, fmt="percent")

    # ---- body & mind block
    week_header(R_METRIC_HDR, "BODY & MIND")
    day_rows(R_METRIC_HDR + 1, R_METRIC_HDR + 2, src=(6, 7))
    put(ws, f"B{R_METRIC_HDR + 1}:D{R_METRIC_HDR + 1}", "Log the numbers you have. Blank is fine.", size=7,
        color=MUTED, bg=PANEL, indent=1)
    put(ws, f"B{R_METRIC_HDR + 2}", "Metric", size=7, bold=True, color=MUTED, bg=PANEL, indent=1)
    put(ws, f"C{R_METRIC_HDR + 2}", "Unit", size=7, bold=True, color=MUTED, bg=PANEL)
    put(ws, f"D{R_METRIC_HDR + 2}", "Target", size=7, bold=True, color=MUTED, bg=PANEL)
    for col, t in zip(("AK", "AL", "AM", "AN"), ("Avg", "Target", "Progress", "Hit %")):
        put(ws, f"{col}{R_METRIC_HDR + 1}", t, size=7, bold=True, color=MUTED, bg=PANEL, h="center")
    paint(ws, f"AK{R_METRIC_HDR + 2}:AN{R_METRIC_HDR + 2}", PANEL)
    put(ws, f"AK{R_METRIC_HDR + 2}:AN{R_METRIC_HDR + 2}", "hit % = days on target ÷ days logged", size=6,
        color=DIM, bg=PANEL, h="center")
    for key, label, unit, fmt, tgt in METRICS:
        r = MROW[key]
        put(ws, f"B{r}", label, size=8, bg=PANEL, indent=1)
        put(ws, f"C{r}", unit, size=6, color=MUTED, bg=PANEL)
        put(ws, f"D{r}", tgt, size=6, color=MUTED, bg=PANEL)
        for j in range(1, 32):
            put(ws, f"{dcol(j)}{r}", size=6 if key in ("steps", "kcal") else 7, color=TEXT, bg=PANEL2, h="center",
                fmt=fmt, border=TILE)
        rng = f"E{r}:AI{r}"
        avg = f'=IFERROR(AVERAGE({rng}),"")'
        hit_from = lambda comp: (f'=IFERROR(SUMIFS({DR(comp)},{DR("month")},{k + 1})/'
                                 f'COUNTIFS({DR("month")},{k + 1},{DR(key)},">0"),"")')
        if key == "weight":
            ak = f'=IF($AP$11=0,"",INDEX({DR("w7")},MATCH($AP$11,{DR("date")},0)))'
            first = f'INDEX({DR("w7")},MATCH($AP$12,{DR("date")},0))'
            put(ws, f"AK{r}", ak, size=8, bold=True, bg=PANEL, h="center", fmt="0.0")
            put(ws, f"AL{r}", "7d avg", size=7, color=MUTED, bg=PANEL, h="center")
            put(ws, f"AM{r}:AN{r}", f'=IF($AP$11=0,"",IFERROR("Δ month "&TEXT(AK{r}-{first},"+0.0;-0.0")&" kg",""))',
                size=7, color=TEAL, bg=PANEL, h="center")
            continue
        put(ws, f"AK{r}", avg, size=8, bold=True, bg=PANEL, h="center",
            fmt={"steps": "#,##0", "protein": "0", "kcal": "#,##0"}.get(key, "0.0"))
        if key == "steps":
            put(ws, f"AL{r}", f'=IFERROR(MAX(E{R_STEPT}:AI{R_STEPT})*1000,"")', size=7, color=MUTED, bg=PANEL, h="center",
                fmt="#,##0")
            put(ws, f"AN{r}", hit_from("c_steps"), size=8, bold=True, bg=PANEL, h="center", fmt="0%")
        elif key == "protein":
            put(ws, f"AL{r}", "=ProtTarget", size=7, color=MUTED, bg=PANEL, h="center", fmt="0")
            put(ws, f"AN{r}", hit_from("c_prot"), size=8, bold=True, bg=PANEL, h="center", fmt="0%")
        elif key == "water":
            put(ws, f"AL{r}", "=WaterHit", size=7, color=MUTED, bg=PANEL, h="center", fmt="0.0")
            put(ws, f"AN{r}", hit_from("c_water"), size=8, bold=True, bg=PANEL, h="center", fmt="0%")
        elif key == "kcal":
            put(ws, f"AL{r}", '=TEXT(KcalMin,"#,##0")&"–"&TEXT(KcalMax,"#,##0")', size=6, color=MUTED, bg=PANEL,
                h="center")
            put(ws, f"AN{r}", f'=IFERROR(COUNTIFS({rng},">="&$AP$7,{rng},"<="&$AP$8)/COUNT({rng}),"")', size=8,
                bold=True, bg=PANEL, h="center", fmt="0%")
        elif key == "sleeph":
            put(ws, f"AL{r}", "=SleepTarget", size=7, color=MUTED, bg=PANEL, h="center", fmt="0.0")
            put(ws, f"AN{r}", f'=IFERROR(COUNTIF({rng},">="&$AP$4)/COUNT({rng}),"")', size=8, bold=True, bg=PANEL,
                h="center", fmt="0%")
        else:
            put(ws, f"AL{r}", "/ 10", size=7, color=MUTED, bg=PANEL, h="center")
            put(ws, f"AN{r}", None, bg=PANEL)
        if key in ("energy", "stress", "mood"):
            put(ws, f"AM{r}", f'=IF(ISNUMBER(AK{r}),{bar(f"AK{r}/10", 10)},"")', size=8,
                color={"energy": YELLOW, "stress": PINK, "mood": TEAL}[key], bg=PANEL)
        elif key == "kcal":
            put(ws, f"AM{r}", f"={bar(f'AN{r}', 10)}", size=8, color=ORANGE, bg=PANEL)
        else:
            put(ws, f"AM{r}", f'=IF(AND(ISNUMBER(AK{r}),ISNUMBER(AL{r})),{bar(f"AK{r}/AL{r}", 10)},"")', size=8,
                color=TEAL, bg=PANEL)
        if key in ("energy", "stress", "mood"):
            dvn = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True,
                                 showErrorMessage=True, errorTitle=label, error="Use a number from 1 to 10.")
            ws.add_data_validation(dvn)
            dvn.add(rng)
    put(ws, f"B{R_W7}", "↳ Weight · 7-day average", size=7, italic=True, color=MUTED, bg=PANEL, indent=2)
    put(ws, f"C{R_W7}:D{R_W7}", "decide on this", size=6, color=DIM, bg=PANEL)
    put(ws, f"B{R_STEPT}", "↳ Steps target (adaptive)", size=7, italic=True, color=MUTED, bg=PANEL, indent=2)
    put(ws, f"C{R_STEPT}:D{R_STEPT}", "7k → 8k → 9k", size=6, color=DIM, bg=PANEL)
    for j in range(1, 32):
        c, r = dcol(j), drow(k, j)
        put(ws, f"{c}{R_W7}", f"=Data!{DC['w7']}{r}", size=6, italic=True, color=MUTED, bg=PANEL, h="center",
            fmt="0.0", border=TILE)
        put(ws, f"{c}{R_STEPT}", f'=IF(Data!{DC["date"]}{r}="","",Data!{DC["stepT"]}{r}/1000)', size=6, italic=True,
            color=MUTED, bg=PANEL, h="center", fmt='0"k"', border=TILE)
    paint(ws, f"AK{R_W7}:AN{R_STEPT}", PANEL)

    # ---- charts: mind & recovery, weight
    r0 = R_STEPT + 2  # 71
    put(ws, f"B{r0}:D{r0}", "TRENDS", size=9, bold=True, color=TEXT, bg=PANEL, indent=1)
    put(ws, f"B{r0 + 1}:D{r0 + 13}",
        "Energy, stress and mood show how the month feels, not just what got done.\n\n"
        "Weight: the dots jump around with water and salt. The 7-day average is the only line that "
        "decides anything.\n\nTarget pace: −0.4 to −0.8 kg a week.",
        size=7, color=MUTED, bg=PANEL, wrap=True, v="top", indent=1)
    put(ws, f"E{r0}:T{r0}", "MIND & RECOVERY · energy · stress · mood", size=8, bold=True, color=TEXT, bg=PANEL,
        indent=1)
    put(ws, f"U{r0}:AI{r0}", "WEIGHT · daily vs 7-day average", size=8, bold=True, color=TEXT, bg=PANEL, indent=1)
    paint(ws, f"E{r0 + 1}:AI{r0 + 13}", PANEL)
    lc = LineChart()
    for key, col in (("en_na", YELLOW), ("st_na", PINK), ("mo_na", TEAL)):
        s = Series(Reference(data_ws, min_col=DATA_KEYS.index(key) + 1, min_row=a, max_row=b),
                   title={"en_na": "Energy", "st_na": "Stress", "mo_na": "Mood"}[key])
        lc.series.append(s)
        line_series(lc, col)
    lc.set_categories(Reference(data_ws, min_col=DATA_KEYS.index("day") + 1, min_row=a, max_row=b))
    style_chart(lc, ymin=0, ymax=10, major=2)
    size_chart(lc, ws, 5, 20, r0 + 1, r0 + 13)
    ws.add_chart(lc)
    chart_style(sn, lc, "line", [YELLOW, PINK, TEAL], vmin=0, vmax=10)
    wc = LineChart()
    s = Series(Reference(data_ws, min_col=DATA_KEYS.index("w_na") + 1, min_row=a, max_row=b), title="Daily")
    wc.series.append(s)
    line_series(wc, mix(VIOLET, PANEL, .7), width=12000, smooth=False)
    wc.series[-1].marker.symbol = "circle"
    wc.series[-1].marker.size = 4
    wc.series[-1].marker.graphicalProperties = GraphicalProperties(solidFill=VIOLET, ln=LineProperties(noFill=True))
    wc.series[-1].graphicalProperties.line.noFill = True
    s = Series(Reference(data_ws, min_col=DATA_KEYS.index("w7_na") + 1, min_row=a, max_row=b), title="7-day avg")
    wc.series.append(s)
    line_series(wc, TEAL, width=28000)
    wc.set_categories(Reference(data_ws, min_col=DATA_KEYS.index("day") + 1, min_row=a, max_row=b))
    style_chart(wc, yfmt="0.0")
    size_chart(wc, ws, 21, MC_LAST, r0 + 1, r0 + 13)
    ws.add_chart(wc)
    chart_style(sn, wc, "weight", [VIOLET, TEAL])

    # ---- conditional formatting (first matching rule wins in Google Sheets)
    full = f"E6:AI{R_STEPT}"
    cf(ws, full, 'E$7=""', fill=BG, color=BG, stop=True)
    for rng, dr in ((f"E6:AI7", 7), (f"E{R_METRIC_HDR + 1}:AI{R_METRIC_HDR + 2}", R_METRIC_HDR + 2)):
        cf(ws, rng, f'AND(E${dr}<>"",E${dr}=$AP$6)', fill=YELLOW, color=BG, bold=True, stop=True)
    for code, col in (("M", YELLOW), ("S", PINK), ("R", TEAL)):
        cf(ws, f"E{R_MODE}:AI{R_MODE}", f'UPPER(E{R_MODE})="{code}"', fill=col, color=BG, bold=True, stop=True)
    for w, (a1, b1) in enumerate(WEEK_BLOCKS):
        for ra, rb in hr:
            top = f"{dcol(a1)}{ra}"
            cf(ws, f"{dcol(a1)}{ra}:{dcol(b1)}{rb}", f'OR({top}=TRUE,{top}="x",{top}="X")',
               fill=mix(WEEK_COLORS[w], PANEL2, .28), stop=True)
    band_cf(ws, f"E{R_SCORE}:AI{R_SCORE}", f"E{R_SCORE}")
    band_cf(ws, f"AM10:AN{R_SCORE}", "$AN10")
    for w in range(5):
        band_cf(ws, f"AN{44 + w * 2}", f"AN{44 + w * 2}")
    band_cf(ws, f"AN{R_METRIC0 + 1}:AN{R_METRIC0 + len(METRICS) - 1}", f"AN{R_METRIC0 + 1}")
    pr, wr, sr, kr = MROW["protein"], MROW["water"], MROW["sleeph"], MROW["kcal"]
    good, bad = mix(GREEN, PANEL2, .35), mix(PINK, PANEL2, .35)
    cf(ws, f"E{pr}:AI{pr}", f'AND(ISNUMBER(E{pr}),E{pr}>=$AP$2)', fill=good, stop=True)
    cf(ws, f"E{wr}:AI{wr}", f'AND(ISNUMBER(E{wr}),E{wr}>=$AP$3)', fill=good, stop=True)
    stp = MROW["steps"]
    cf(ws, f"E{stp}:AI{stp}", f'AND(ISNUMBER(E{stp}),ISNUMBER(E${R_STEPT}),E{stp}>=E${R_STEPT}*1000)', fill=good,
       stop=True)
    cf(ws, f"E{sr}:AI{sr}", f'AND(ISNUMBER(E{sr}),E{sr}>=$AP$4)', fill=good, stop=True)
    cf(ws, f"E{sr}:AI{sr}", f'AND(ISNUMBER(E{sr}),E{sr}<$AP$5)', fill=bad, stop=True)
    cf(ws, f"E{kr}:AI{kr}", f'AND(ISNUMBER(E{kr}),E{kr}>=$AP$7,E{kr}<=$AP$8)', fill=good, stop=True)
    for key in ("energy", "mood"):
        r = MROW[key]
        cf(ws, f"E{r}:AI{r}", f'AND(ISNUMBER(E{r}),E{r}>=7)', fill=good, stop=True)
        cf(ws, f"E{r}:AI{r}", f'AND(ISNUMBER(E{r}),E{r}<=3)', fill=bad, stop=True)
    r = MROW["stress"]
    cf(ws, f"E{r}:AI{r}", f'AND(ISNUMBER(E{r}),E{r}>=8)', fill=bad, stop=True)
    cf(ws, f"E{r}:AI{r}", f'AND(ISNUMBER(E{r}),E{r}<=3)', fill=good, stop=True)

    ws.freeze_panes = "E8"
    ws.column_dimensions["AP"].hidden = False


# ======================================================================== DATA ENGINE
def build_data(wb, ws):
    ws.sheet_properties.tabColor = "39425C"
    page(ws)
    ws["A1"] = "DATA ENGINE · calculated from the month tabs · don't type here"
    ws["A1"].font = F(11, YELLOW, True)
    ws["A2"] = ("One row per calendar day. Habit columns are 1/0, metrics are your numbers, then the scoring. "
                "Rows for dates that don't exist (e.g. 31 Nov) stay blank.")
    ws["A2"].font = F(8, MUTED)
    for i, key in enumerate(DATA_KEYS):
        c = ws.cell(row=3, column=i + 1, value=key)
        c.font = F(7, TEXT, True)
        c.fill = PF(PANEL2)
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[CL(i + 1)].width = 7 if i else 11
    ws.freeze_panes = "B4"
    d = DC
    s_rng, d_rng, w_rng = DR("score"), DR("date"), DR("weight")
    l1 = [d[k] for k in L1KEYS]
    for k, y, m, sn in MONTHS:
        S = f"'{sn}'!"
        for j in range(1, 32):
            r = drow(k, j)
            col = dcol(j)
            A = f"$A{r}"
            f = {}
            f["date"] = f'=IF({S}{col}7="","",{S}{col}7)'
            f["month"] = k + 1
            f["day"] = f'=IF({A}="","",{j})'
            f["wd"] = f'=IF({A}="","",WEEKDAY({A},2))'
            f["mode"] = f'=IF({A}="","",UPPER(TRIM({S}{col}{R_MODE})))'
            for hk in HKEYS:
                ref = f"{S}{col}{HROW[hk]}"
                f[hk] = f'=IF({A}="","",IF(OR({ref}=TRUE,{ref}="x",{ref}="X"),1,0))'
            for mk in MKEYS:
                ref = f"{S}{col}{MROW[mk]}"
                f[mk] = f'=IF({A}="","",IF({ref}="","",{ref}))'
            g = lambda key: f"{d[key]}{r}"
            f["active"] = f'=IF({A}="",0,IF(AND({A}>=TrackStart,{A}<={TODAY}),1,0))'
            f["logged"] = (f'=IF({A}="",0,IF(SUM({d[HKEYS[0]]}{r}:{d[HKEYS[-1]]}{r})+'
                           f'COUNT({d[MKEYS[0]]}{r}:{d[MKEYS[-1]]}{r})>0,1,0))')
            f["wk"] = f'=IF({A}="","",INT(({A}-TrackStart)/7)+1)'
            f["stepT"] = f'=IF({A}="","",IF({g("wk")}<=1,Steps1,IF({g("wk")}=2,Steps2,Steps3)))'
            f["workday"] = f'=IF({A}="",0,IF(ISNUMBER(SEARCH(TEXT({g("wd")},"0"),WorkDays&"")),1,0))'
            f["strday"] = f'=IF({A}="",0,IF(ISNUMBER(SEARCH(TEXT({g("wd")},"0"),StrDays&"")),1,0))'
            f["c_sleep"] = (f'=IF({A}="",0,IF(ISNUMBER({g("sleeph")}),IF({g("sleeph")}>=SleepMin,1,0),'
                            f'{g("sleep")}))')
            f["c_move"] = f'=IF({A}="",0,MAX({g("move")},{g("str")},{g("cardio")},{g("xcardio")}))'
            f["c_prot"] = f'=IF({A}="",0,IF(ISNUMBER({g("protein")}),IF({g("protein")}>=ProtHit,1,0),0))'
            f["c_water"] = f'=IF({A}="",0,IF(ISNUMBER({g("water")}),IF({g("water")}>=WaterHit,1,0),0))'
            f["c_steps"] = f'=IF({A}="",0,IF(ISNUMBER({g("steps")}),IF({g("steps")}>={g("stepT")},1,0),0))'
            f["c_pray"] = f'=IF({A}="",0,{g("pray")})'
            f["c_work"] = f'=IF({A}="",0,{g("top3")})'
            f["c_love"] = f'=IF({A}="",0,{g("love")})'
            f["c_read"] = f'=IF({A}="",0,MAX({g("read")},{g("xread")}))'
            f["norm"] = (f'=IF({A}="","",({g("c_sleep")}+{g("c_move")}+{g("c_prot")}+{g("c_water")}+{g("c_steps")}+'
                         f'{g("c_pray")}+{g("c_love")}+{g("c_read")}+IF({g("workday")}=1,{g("c_work")},0))'
                         f'/(8+{g("workday")}))')
            f["mvd"] = (f'=IF({A}="","",(IF(ISNUMBER({g("water")}),IF({g("water")}>=WaterMVD,1,0),0)+{g("c_move")}+'
                        f'IF(ISNUMBER({g("protein")}),IF({g("protein")}>=ProtMVD,1,0),0)+'
                        f'MAX({g("mob")},{g("xmob")})+{g("c_pray")}+{g("c_love")}+{g("c_sleep")})/7)')
            f["score"] = (f'=IF({g("active")}=0,"",IF({g("mode")}="S","",IF(AND({g("logged")}=0,{A}={TODAY}),"",'
                          f'IF({g("mode")}="M",{g("mvd")},{g("norm")}))))')
            f["r7"] = (f'=IF({g("active")}=0,"",IFERROR(AVERAGEIFS({s_rng},{d_rng},">"&({A}-7),{d_rng},"<="&{A}),""))')
            f["r30"] = (f'=IF({g("active")}=0,"",IFERROR(AVERAGEIFS({s_rng},{d_rng},">"&({A}-30),{d_rng},"<="&{A}),""))')
            f["w7"] = (f'=IF({A}="","",IFERROR(AVERAGEIFS({w_rng},{d_rng},">"&({A}-7),{d_rng},"<="&{A},'
                       f'{w_rng},">0"),""))')
            f["l1done"] = f'=IF({A}="","",{"+".join(f"{c}{r}" for c in l1)}-IF({g("workday")}=1,0,{g("top3")}))'
            f["l1app"] = f'=IF({A}="","",8+{g("workday")})'
            f["score_na"] = f'=IF(ISNUMBER({g("score")}),{g("score")},NA())'
            f["w_na"] = f'=IF(ISNUMBER({g("weight")}),{g("weight")},NA())'
            f["w7_na"] = f'=IF(ISNUMBER({g("w7")}),{g("w7")},NA())'
            f["en_na"] = f'=IF(ISNUMBER({g("energy")}),{g("energy")},NA())'
            f["st_na"] = f'=IF(ISNUMBER({g("stress")}),{g("stress")},NA())'
            f["mo_na"] = f'=IF(ISNUMBER({g("mood")}),{g("mood")},NA())'
            for key in DATA_KEYS:
                c = ws[f"{d[key]}{r}"]
                c.value = f[key]
                c.font = F(7, TEXT)
            ws[f"A{r}"].number_format = "ddd d mmm yy"
            for key in ("score", "norm", "mvd", "r7", "r30"):
                ws[f"{d[key]}{r}"].number_format = "0%"
            ws[f"{d['w7']}{r}"].number_format = "0.0"
    ws.column_dimensions["A"].width = 13


# ======================================================================== CALC (dashboard engine)
CALC = {}


def build_calc(wb, ws):
    ws.sheet_properties.tabColor = "39425C"
    page(ws)
    widths(ws, {"A": 22, "B": 14, "C": 3, "D": 11, "E": 7, "F": 7, "G": 7, "H": 7, "I": 7, "J": 7, "K": 7, "L": 7,
                "M": 7, "N": 7, "O": 3, "P": 11, "Q": 8, "R": 3, "S": 11, "T": 8})
    ws["A1"] = "CALC · dashboard engine · don't type here"
    ws["A1"].font = F(11, YELLOW, True)

    def val(key, dexpr="AsOf"):
        return f'IFERROR(INDEX({DR(key)},MATCH({dexpr},{DR("date")},0)),"")'

    def win(days, end="AsOf"):
        return f'{DR("date")},">"&({end}-{days}),{DR("date")},"<="&{end}'

    def sumw(key, days):
        return f"SUMIFS({DR(key)},{win(days)})"

    def avgw(key, days):
        return f'IFERROR(AVERAGEIFS({DR(key)},{win(days)},{DR(key)},">0"),"")'

    lib = "'Habit Library'!"
    rows = [
        ("k_Mode", val("mode"), "@"),
        ("k_Wd", "=WEEKDAY(AsOf,2)", "0"),
        ("k_PlanT", "=INDEX(PlanTraining,k_Wd)", "@"),
        ("k_PlanR", "=INDEX(PlanRest,k_Wd)", "@"),
    ]
    for hk in HKEYS + MKEYS + ["stepT", "score", "r7", "r30", "w7", "l1done", "l1app", "workday", "strday",
                               "logged", "active"]:
        rows.append((f"k_T_{hk}", val(hk), "General"))
    rows += [
        ("k_Y_score", val("score", "AsOf-1"), "0%"),
        ("k_Y_str", val("str", "AsOf-1"), "0"),
        ("k_Y_strday", val("strday", "AsOf-1"), "0"),
        ("k_Y_mode", val("mode", "AsOf-1"), "@"),
        ("k_Y_active", val("active", "AsOf-1"), "0"),
        ("k_W7_m7", val("w7", "AsOf-7"), "0.0"),
        ("k_W7_m14", val("w7", "AsOf-14"), "0.0"),
        ("k_W7_m30", val("w7", "AsOf-30"), "0.0"),
        ("k_R7_m7", val("r7", "AsOf-7"), "0%"),
        ("k_Rate", '=IF(OR(NOT(ISNUMBER(k_T_w7)),NOT(ISNUMBER(k_W7_m7))),"",k_T_w7-k_W7_m7)', "+0.00;-0.00"),
        ("k_RatePrev", '=IF(OR(NOT(ISNUMBER(k_W7_m7)),NOT(ISNUMBER(k_W7_m14))),"",k_W7_m7-k_W7_m14)', "+0.00;-0.00"),
        ("k_Change30", '=IF(OR(NOT(ISNUMBER(k_T_w7)),NOT(ISNUMBER(k_W7_m30))),"",k_T_w7-k_W7_m30)', "+0.0;-0.0"),
        ("k_TodayOpen", '=IF(AND(N(k_T_active)=1,N(k_T_logged)=0,k_Mode<>"S"),1,0)', "0"),
        ("k_D7", f'=COUNTIFS({win(7)},{DR("active")},1,{DR("mode")},"<>S")-k_TodayOpen', "0"),
        ("k_D30", f'=COUNTIFS({win(30)},{DR("active")},1,{DR("mode")},"<>S")-k_TodayOpen', "0"),
        ("k_WD7", f'=COUNTIFS({win(7)},{DR("active")},1,{DR("mode")},"<>S",{DR("workday")},1)-IF(N(k_T_workday)=1,k_TodayOpen,0)', "0"),
        ("k_WD30", f'=COUNTIFS({win(30)},{DR("active")},1,{DR("mode")},"<>S",{DR("workday")},1)-IF(N(k_T_workday)=1,k_TodayOpen,0)', "0"),
        ("k_Logged7", "=" + sumw("logged", 7), "0"),
        ("k_Str7", "=" + sumw("str", 7), "0"),
        ("k_Str14", "=" + sumw("str", 14), "0"),
        ("k_Prog14", "=" + sumw("prog", 14), "0"),
        ("k_Str30", "=" + sumw("str", 30), "0"),
        ("k_ProgRate14", '=IF(k_Str14=0,"",MIN(1,k_Prog14/k_Str14))', "0%"),
        ("k_Core7", "=" + sumw("core", 7), "0"),
        ("k_Calf7", "=" + sumw("calf", 7), "0"),
        ("k_Cardio7", f'={sumw("cardio", 7)}+{sumw("xcardio", 7)}', "0"),
        ("k_Church7", "=" + sumw("church", 7), "0"),
        ("k_Son7", "=" + sumw("son", 7), "0"),
        ("k_Date7", "=" + sumw("couple", 7), "0"),
        ("k_Family7", "=" + sumw("family", 7), "0"),
        ("k_Prep7", "=" + sumw("prep", 7), "0"),
        ("k_Pray7", "=" + sumw("pray", 7), "0"),
        ("k_Pray30", "=" + sumw("pray", 30), "0"),
        ("k_Read30", "=" + sumw("c_read", 30), "0"),
        ("k_Prot7", "=" + avgw("protein", 7), "0"),
        ("k_Steps7", "=" + avgw("steps", 7), "#,##0"),
        ("k_Sleep7", "=" + avgw("sleeph", 7), "0.0"),
        ("k_Energy7", "=" + avgw("energy", 7), "0.0"),
        ("k_Stress7", "=" + avgw("stress", 7), "0.0"),
        ("k_Mood7", "=" + avgw("mood", 7), "0.0"),
        ("k_Water7", "=" + avgw("water", 7), "0.0"),
        ("k_Rest7", f'=COUNTIFS({win(7)},{DR("active")},1,{DR("str")},0)', "0"),
        ("k_Recovery", '=IF(k_Logged7=0,"",ROUND(100*(0.3*IF(ISNUMBER(k_Sleep7),MIN(1,k_Sleep7/SleepTarget),0.7)+'
                       '0.2*IF(ISNUMBER(k_Energy7),k_Energy7/10,0.7)+0.2*IF(ISNUMBER(k_Stress7),(10-k_Stress7)/9,0.7)+'
                       '0.15*IF(k_Rest7>=1,1,0)+0.15*IF(k_Str14=0,0.5,MIN(1,k_ProgRate14/0.5))),0))', "0"),
        ("k_L1avg7", f'=IFERROR(AVERAGEIFS({lib}$K$6:$K$40,{lib}$C$6:$C$40,"L1"),"")', "0%"),
        ("k_L2avg7", f'=IFERROR(AVERAGEIFS({lib}$K$6:$K$40,{lib}$C$6:$C$40,"L2"),"")', "0%"),
        ("k_L1low", f'=IF(k_D7<4,0,COUNTIFS({lib}$C$6:$C$40,"L1",{lib}$K$6:$K$40,"<0.5"))', "0"),
        ("k_Status7", "=" + status("k_T_r7"), "@"),
        ("k_Status30", "=" + status("k_T_r30"), "@"),
        ("k_WaistDate", "=_xlfn.MAXIFS(Body!$B$17:$B$43,Body!$B$17:$B$43,\"<=\"&AsOf,Body!$D$17:$D$43,\">0\")", "d mmm"),
        ("k_Waist", '=IF(k_WaistDate=0,"",INDEX(Body!$D$17:$D$43,MATCH(k_WaistDate,Body!$B$17:$B$43,0)))', "0.0"),
        ("k_WaistPrevDate", "=_xlfn.MAXIFS(Body!$B$17:$B$43,Body!$B$17:$B$43,\"<\"&k_WaistDate,Body!$D$17:$D$43,\">0\")",
         "d mmm"),
        ("k_WaistPrev", '=IF(k_WaistPrevDate=0,"",INDEX(Body!$D$17:$D$43,MATCH(k_WaistPrevDate,Body!$B$17:$B$43,0)))',
         "0.0"),
        ("k_Photo30", '=COUNTIFS(Body!$B$17:$B$43,">"&(AsOf-31),Body!$B$17:$B$43,"<="&AsOf,Body!$J$17:$J$43,TRUE)+'
                      'COUNTIFS(Body!$B$17:$B$43,">"&(AsOf-31),Body!$B$17:$B$43,"<="&AsOf,Body!$J$17:$J$43,"x")', "0"),
        ("k_WeeksToGoal", '=IF(OR(NOT(ISNUMBER(k_Rate)),NOT(ISNUMBER(k_T_w7))),"",IF(k_Rate>=0,"",'
                          'MAX(0,(k_T_w7-GoalWeight)/(-k_Rate))))', "0"),
        ("k_LossTooFast", '=IF(AND(ISNUMBER(k_Rate),ISNUMBER(k_RatePrev)),IF(AND(-k_Rate>LossMaxPct*k_W7_m7,'
                          '-k_RatePrev>LossMaxPct*k_W7_m14),1,0),0)', "0"),
        ("k_RateFlag", '=IF(NOT(ISNUMBER(k_Rate)),"needs 2 weeks of weigh-ins",IF(-k_Rate>LossMaxPct*k_W7_m7,"faster than 1%/wk",'
                       'IF(-k_Rate>=LossMin,IF(-k_Rate<=LossMax,"on target","fast end"),IF(-k_Rate>0,"slow","flat / up"))))', "@"),
    ]
    put(ws, "A3", "Name", size=8, bold=True, color=MUTED)
    put(ws, "B3", "Value", size=8, bold=True, color=MUTED)
    r = 4
    for nm, f, fmt in rows:
        ws[f"A{r}"] = nm
        ws[f"A{r}"].font = F(7, MUTED)
        c = ws[f"B{r}"]
        c.value = f if str(f).startswith("=") else "=" + f
        c.font = F(8, TEXT)
        if fmt != "General":
            c.number_format = fmt
        name(wb, nm, f"Calc!$B${r}")
        CALC[nm] = r
        r += 1

    # ---- dashboard series (60 days ending AsOf)
    hdr = ["date", "score", "r7", "Daily", "7-day avg", "steps", "target", "sleep", "energy", "stress", "mood"]
    for i, hd in enumerate(hdr):
        put(ws, f"{CL(4 + i)}3", hd, size=7, bold=True, color=MUTED)
    for i in range(60):
        rr = 4 + i
        dexpr = f"$D{rr}"
        ws[f"D{rr}"] = f"=AsOf-{59 - i}"
        ws[f"D{rr}"].number_format = "d mmm"

        def na(key):
            return f'=IFERROR(IF(ISNUMBER({val(key, dexpr)}),{val(key, dexpr)},NA()),NA())'

        ws[f"E{rr}"] = na("score")
        ws[f"F{rr}"] = na("r7")
        ws[f"G{rr}"] = na("weight")
        ws[f"H{rr}"] = na("w7")
        ws[f"I{rr}"] = na("steps")
        ws[f"J{rr}"] = f'=IFERROR(IF(ISNUMBER({val("stepT", dexpr)}),{val("stepT", dexpr)},NA()),NA())'
        ws[f"K{rr}"] = na("sleeph")
        ws[f"L{rr}"] = na("energy")
        ws[f"M{rr}"] = na("stress")
        ws[f"N{rr}"] = na("mood")
        for col in "DEFGHIJKLMN":
            ws[f"{col}{rr}"].font = F(7, TEXT)
        for col in "EF":
            ws[f"{col}{rr}"].number_format = "0%"
    # ---- weekly weight + waist series for the Body tab
    put(ws, "P3", "week", size=7, bold=True, color=MUTED)
    put(ws, "Q3", "avg kg", size=7, bold=True, color=MUTED)
    for i in range(53):
        rr = 4 + i
        ws[f"P{rr}"] = f"=DATE({FIRST_MONDAY.year},{FIRST_MONDAY.month},{FIRST_MONDAY.day})+{7 * i}"
        ws[f"P{rr}"].number_format = "d mmm"
        ws[f"Q{rr}"] = (f'=IFERROR(AVERAGEIFS({DR("weight")},{DR("date")},">="&P{rr},{DR("date")},"<="&(P{rr}+6),'
                        f'{DR("weight")},">0"),NA())')
        ws[f"Q{rr}"].number_format = "0.0"
    put(ws, "S3", "measured", size=7, bold=True, color=MUTED)
    put(ws, "T3", "waist", size=7, bold=True, color=MUTED)
    for i in range(27):
        rr = 4 + i
        ws[f"S{rr}"] = f"=Body!B{17 + i}"
        ws[f"S{rr}"].number_format = "d mmm"
        ws[f"T{rr}"] = f'=IF(ISNUMBER(Body!D{17 + i}),Body!D{17 + i},NA())'


# ======================================================================== HABIT LIBRARY
LIB_FIRST = 6


def build_library(wb, ws):
    page(ws)
    ws.sheet_properties.tabColor = TEAL
    paint(ws, "A1:P45", BG)
    widths(ws, {"A": 2, "B": 26, "C": 6, "D": 13, "E": 12, "F": 12, "G": 7, "H": 62, "I": 5, "J": 7, "K": 8,
                "L": 7, "M": 8, "N": 14, "O": 7})
    put(ws, "B2", "HABIT LIBRARY", size=22, bold=True)
    put(ws, "B3", '="The master list behind every tab · scores as of "&TEXT(AsOf,"d mmm yyyy")&'
                  '" · score = done ÷ (weekly target × days tracked)"', size=8, color=MUTED)
    hdr = ["Habit", "Level", "Pillar", "Frequency", "Target", "/ week", "What counts", "MVD", "7d done",
           "7d score", "30d done", "30d score", "30-day consistency"]
    for i, hd in enumerate(hdr):
        put(ws, f"{CL(2 + i)}5", hd, size=7, bold=True, color=MUTED, bg=PANEL2, h="left" if i in (0, 6, 12) else "center",
            indent=1 if i in (0, 6) else 0)
    items = [(h[0], h[1], h[2], h[3], h[4], h[5], h[6], h[7], "◆" if h[8] else "") for h in HABITS]
    items += [(a[0], a[1], a[2], a[3], a[4], a[5], a[6], a[7] + "  (auto)", "◆" if a[0] in ("c_water", "c_prot", "c_sleep") else "")
              for a in AUTO]
    r = LIB_FIRST
    for key, nm, lvl, pillar, freq, target, weekly, what, mvd in items:
        bg = PANEL if r % 2 else PANEL2
        put(ws, f"B{r}", nm, size=8, bg=bg, indent=1)
        put(ws, f"C{r}", lvl, size=8, bold=True, color=LEVEL_COLORS[lvl], bg=bg, h="center")
        put(ws, f"D{r}", pillar, size=8, bold=True, color=PILLARS[pillar], bg=bg, h="center")
        put(ws, f"E{r}", freq, size=8, color=MUTED, bg=bg, h="center")
        put(ws, f"F{r}", target, size=8, color=MUTED, bg=bg, h="center")
        put(ws, f"G{r}", weekly, size=8, bg=bg, h="center")
        put(ws, f"H{r}", what, size=7, color=MUTED, bg=bg, wrap=True, indent=1)
        put(ws, f"I{r}", mvd, size=8, color=YELLOW, bg=bg, h="center")
        col = DR(key)
        put(ws, f"J{r}", f'=SUMIFS({col},{DR("date")},">"&(AsOf-7),{DR("date")},"<="&AsOf)', size=8, bg=bg, h="center")
        put(ws, f"L{r}", f'=SUMIFS({col},{DR("date")},">"&(AsOf-30),{DR("date")},"<="&AsOf)', size=8, bg=bg, h="center")
        if weekly == 0:
            put(ws, f"K{r}", "bonus", size=7, color=DIM, bg=bg, h="center")
            put(ws, f"M{r}", "bonus", size=7, color=DIM, bg=bg, h="center")
            put(ws, f"N{r}:O{r}", f'=IF(L{r}=0,"",L{r}&" × in 30 days")', size=7, color=MUTED, bg=bg)
        else:
            e7 = "k_WD7" if key == "top3" else f"k_D7*{weekly}/7"
            e30 = "k_WD30" if key == "top3" else f"k_D30*{weekly}/7"
            put(ws, f"K{r}", f'=IFERROR(MIN(1,J{r}/({e7})),"")', size=8, bold=True, bg=bg, h="center", fmt="0%")
            put(ws, f"M{r}", f'=IFERROR(MIN(1,L{r}/({e30})),"")', size=8, bold=True, bg=bg, h="center", fmt="0%")
            put(ws, f"N{r}:O{r}", f"={bar(f'M{r}', 10)}", size=8, color=TEAL, bg=bg)
        ws.row_dimensions[r].height = 26
        r += 1
    band_cf(ws, f"K{LIB_FIRST}:K{r}", f"K{LIB_FIRST}")
    band_cf(ws, f"M{LIB_FIRST}:M{r}", f"M{LIB_FIRST}")
    band_cf(ws, f"N{LIB_FIRST}:N{r}", f"M{LIB_FIRST}")
    put(ws, f"B{r + 1}",
        "◆ = part of the Minimum Viable Day. L3 habits never count against you. Weekly items are scored against "
        "their weekly target, so one church visit a week is 100%.", size=7, color=MUTED, italic=True)
    ws.row_dimensions[r + 1].height = 26
    ws.freeze_panes = "C6"
    return r - 1


# ======================================================================== DASHBOARD
def build_dashboard(wb, ws):
    page(ws)
    ws.sheet_properties.tabColor = YELLOW
    paint(ws, "A1:T86", BG)
    widths(ws, {"A": 2, **{CL(c): 9.5 for c in range(2, 20)}, "T": 2})
    heights(ws, range(1, 87), 18)

    put(ws, "B2:J3", "LIFE OS", size=28, bold=True, v="center")
    put(ws, "B4:M4", "Lucas · personal operating system  ·  Consistency > Intensity  ·  Progress > Perfection  ·  "
                     "Systems > Motivation  ·  Health > Extreme results", size=7, color=MUTED)
    put(ws, "N2:O2", "SHOWING", size=7, bold=True, color=MUTED, bg=PANEL, h="right")
    c = put(ws, "P2:S2", f"={TODAY}", size=10, bold=True, color=YELLOW, bg=PANEL2, h="center",
            fmt="dddd d mmmm yyyy")
    name(wb, "AsOf", "Dashboard!$P$2")
    put(ws, "N3:O3", "DAY MODE", size=7, bold=True, color=MUTED, bg=PANEL, h="right")
    put(ws, "P3:S3", '=IF(k_Mode="M","Minimum Viable Day",IF(k_Mode="S","Sick day · score paused",'
                     'IF(k_Mode="R","Rest / recovery day","Normal day")))', size=8, bold=True, bg=PANEL2, h="center")
    put(ws, "N4:S4", "type any date in the yellow cell to look back · =TODAY() for live", size=6, color=DIM, h="right")
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 22

    # KPI cards
    cards = [
        ("B", "D", "FOUNDATION · TODAY", "=k_T_score", "0%",
         '=IF(NOT(ISNUMBER(k_T_score)),IF(k_Mode="S","sick day · paused","not logged yet"),'
         'k_T_l1done&" of "&k_T_l1app&" non-negotiables")', YELLOW),
        ("E", "G", "ROLLING 7-DAY", "=k_T_r7", "0%", "=k_Status7", TEAL),
        ("H", "J", "ROLLING 30-DAY", "=k_T_r30", "0%", "=k_Status30", TEAL),
        ("K", "M", "STATUS · 7 DAY", "=k_Status7", "@", '="Green ≥85 · Good 70–84 · Attention 50–69 · Reset <50"', TEXT),
        ("N", "P", "WEIGHT · 7-DAY AVG", '=IF(ISNUMBER(k_T_w7),k_T_w7,"—")', '0.0" kg"',
         '=IF(ISNUMBER(k_T_weight),"today "&TEXT(k_T_weight,"0.0")&" kg","no weigh-in today")', VIOLET),
        ("Q", "S", "WEEKLY RATE", '=IF(ISNUMBER(k_Rate),k_Rate,"—")', '+0.00" kg";-0.00" kg"', "=k_RateFlag", GREEN),
    ]
    for c1, c2, label, f, fmt, sub, col in cards:
        put(ws, f"{c1}6:{c2}6", label, size=7, bold=True, color=MUTED, bg=PANEL, h="center", v="bottom")
        put(ws, f"{c1}7:{c2}8", f, size=24 if fmt != "@" else 14, bold=True, color=col, bg=PANEL, h="center", fmt=fmt)
        put(ws, f"{c1}9:{c2}9", sub, size=7 if c1 != "K" else 6, color=MUTED, bg=PANEL, h="center", v="top")
    status_cf(ws, "K7:M8", "$K$7")
    status_cf(ws, "E9:J9", "E9")
    band_cf(ws, "B7:J8", "B7")

    put(ws, "B11:S11", '="TODAY  ·  "&UPPER(TEXT(AsOf,"dddd"))&"   ·   06:00 wake + reset   ·   06:30 "&k_PlanT&'
                       '"   ·   10:00–20:00 work   ·   "&k_PlanR&"   ·   22:00 lights out"',
        size=9, bold=True, color=TEXT, bg=PANEL2, indent=1)
    ws.row_dimensions[11].height = 24

    # TODAY cards
    def tick(key):
        return f'=IF(k_T_{key}=1,"✓ done",IF(k_T_{key}="","—","○ open"))'

    def num(key, fmt, tgt=None, unit=""):
        base = f'IF(ISNUMBER(k_T_{key}),TEXT(k_T_{key},"{fmt}")'
        return f'={base}&"{unit}"' + (f'&" / "&{tgt}' if tgt else "") + ',"—")'

    groups = [
        ("BODY", PINK, "B", "C", "D", "E", [
            ("Weight", num("weight", "0.0", unit=" kg")),
            ("Steps", '=IF(ISNUMBER(k_T_steps),TEXT(k_T_steps,"#,##0")&" / "&TEXT(k_T_stepT,"#,##0"),"— / "&TEXT(k_T_stepT,"#,##0"))'),
            ("Workout", tick("move")),
            ("Protein", '=IF(ISNUMBER(k_T_protein),TEXT(k_T_protein,"0")&" / "&ProtTarget&" g","— / "&ProtTarget&" g")'),
            ("Water", '=IF(ISNUMBER(k_T_water),TEXT(k_T_water,"0.0")&" / "&TEXT(WaterHit,"0.0")&" L","— / "&TEXT(WaterHit,"0.0")&" L")'),
            ("Calories", '=IF(ISNUMBER(k_T_kcal),TEXT(k_T_kcal,"#,##0"),"—")'),
        ]),
        ("MIND", VIOLET, "F", "G", "H", "H", [
            ("Read / learn", tick("read")),
            ("Deep work", tick("deep")),
            ("Med / journal", tick("med")),
            ("Mood", num("mood", "0", unit="/10")),
        ]),
        ("SPIRIT", YELLOW, "I", "J", "K", "K", [
            ("Prayer + Word", tick("pray")),
            ("Church · wk", '=IF(k_Church7>0,"✓ "&k_Church7,"○ 0")'),
            ("Prayer · 7d", '=k_Pray7&" / 7"'),
        ]),
        ("LIFE", FUCHSIA, "L", "M", "N", "O", [
            ("Top 3 set", tick("top3")),
            ("Fiancée time", tick("love")),
            ("Son · this week", '=IF(k_Son7>0,"✓ "&k_Son7,"○ 0")'),
            ("Family · this week", '=IF(k_Family7>0,"✓ "&k_Family7,"○ 0")'),
            ("Home reset", tick("home")),
            ("Shutdown", tick("shut")),
        ]),
        ("RECOVERY", TEAL, "P", "Q", "R", "S", [
            ("Sleep", num("sleeph", "0.0", unit=" h")),
            ("Lights out", tick("sleep")),
            ("Energy", num("energy", "0", unit="/10")),
            ("Stress", num("stress", "0", unit="/10")),
            ("Recovery · 7d", '=IF(ISNUMBER(k_Recovery),k_Recovery&" / 100","—")'),
        ]),
    ]
    put(ws, "B13:S13", "TODAY", size=9, bold=True, color=MUTED)
    for title, col, l1, l2, v1, v2, lines in groups:
        put(ws, f"{l1}14:{v2}14", title, size=9, bold=True, color=col, bg=mix(col, PANEL, .16), indent=1)
        for i in range(6):
            r = 15 + i
            if i < len(lines):
                lab, f = lines[i]
                put(ws, f"{l1}{r}:{l2}{r}" if l1 != l2 else f"{l1}{r}", lab, size=7, color=MUTED, bg=PANEL, indent=1)
                put(ws, f"{v1}{r}:{v2}{r}" if v1 != v2 else f"{v1}{r}", f, size=8, bold=True, bg=PANEL, h="right")
            else:
                paint(ws, f"{l1}{r}:{v2}{r}", PANEL)
        cf(ws, f"{v1}15:{v2}20", f'LEFT({v1}15,1)="✓"', color=GREEN, stop=True)
        cf(ws, f"{v1}15:{v2}20", f'LEFT({v1}15,1)="○"', color=DIM, stop=True)

    # SMART COACH
    put(ws, "B22:L22", "SMART COACH · adaptive recommendations", size=9, bold=True, color=YELLOW, bg=PANEL, indent=1)
    rules = [
        '=IF(k_Mode="S","🤒 Sick-day rule: rest, fluids, protein-conscious food, sleep. No hard training, no deficit, no forced cardio. Resume gradually.","✓ Not a sick day.")',
        '=IF(NOT(ISNUMBER(k_T_sleeph)),"· Log sleep hours to unlock recovery advice.",IF(k_T_sleeph<SleepLow,"⚠ Under "&SleepLow&" h sleep → make today a recovery day: walk + mobility, or cut the session to 2 easy sets per move.","✓ Sleep "&TEXT(k_T_sleeph,"0.0")&" h, train as planned."))',
        '=IF(NOT(ISNUMBER(k_T_protein)),"· Log protein to get a dinner nudge.",IF(k_T_protein<ProtDinner,"⚠ Protein at "&k_T_protein&" g → make dinner high-protein: chicken, eggs, lean beef, dholl + a whey shake (+25 g).","✓ Protein on track ("&k_T_protein&" g)."))',
        f'=IF(AND(AsOf={TODAY},HOUR(NOW())>=18,ISNUMBER(k_T_steps),k_T_steps<StepsAlert),"⚠ Under "&TEXT(StepsAlert,"#,##0")&" steps after 18:00 → 20–30 min walk (with your fiancée or son counts twice).",IF(ISNUMBER(k_T_steps),"✓ Steps "&TEXT(k_T_steps,"#,##0")&" of "&TEXT(k_T_stepT,"#,##0")&" target.","· Steps target today: "&TEXT(k_T_stepT,"#,##0")&"."))',
        '=IF(AND(k_Y_strday=1,k_Y_active=1,k_Y_mode<>"S",k_Y_str=0),"⚠ Yesterday\'s session was missed → don\'t reshuffle the week. Do the 20-minute minimum workout today (Plan & Routines).","✓ Training on plan. Today: "&k_PlanT&".")',
        '=IF(AND(ISNUMBER(k_T_w7),ISNUMBER(k_W7_m14)),IF(k_T_w7>=k_W7_m14,"⚠ 7-day average hasn\'t dropped in 14 days → trim 150–200 kcal or add ~2,000 steps a day. One change only.","✓ Weight trend moving the right way."),"· The weight rules switch on after 2 weeks of weigh-ins.")',
        '=IF(k_LossTooFast=1,"⚠ Losing more than 1% bodyweight a week, twice in a row → add 150–200 kcal and check recovery. Fat loss, not muscle loss.",IF(ISNUMBER(k_Rate),"✓ Rate of loss is sustainable (target 0.4–0.8 kg/week).","· The pace check starts after 2 weeks of weigh-ins."))',
        '=IF(NOT(ISNUMBER(k_Prot7)),"· Protein average appears after a few logged days.",IF(k_Prot7<ProtHit,"⚠ Protein averaging "&TEXT(k_Prot7,"0")&" g this week → muscle-preservation risk. Add one whey serving or an extra egg/chicken portion.","✓ Protein averaging "&TEXT(k_Prot7,"0")&" g."))',
        '=IF(NOT(ISNUMBER(k_Sleep7)),"· Sleep average appears after a few logged nights.",IF(k_Sleep7<SleepMin,"⚠ Sleep averaging "&TEXT(k_Sleep7,"0.0")&" h → protect 22:00 lights out before adding anything else.","✓ Sleep averaging "&TEXT(k_Sleep7,"0.0")&" h."))',
        '=IF(AND(k_Str14>=3,ISNUMBER(k_ProgRate14),k_ProgRate14<0.25),"⚠ Strength is stalling (progress in "&TEXT(k_ProgRate14,"0%")&" of sessions) → check sleep, protein and calories before pushing harder.",IF(ISNUMBER(k_ProgRate14),"✓ Progressive overload: "&TEXT(k_ProgRate14,"0%")&" of sessions improved something.","· Tick Progression achieved after each session to track overload."))',
        '=IF(AND(ISNUMBER(k_Recovery),k_Recovery<RecoveryLow),"⚠ Recovery score "&k_Recovery&" → lighter training this week: one set fewer, stop 2–3 reps before failure.",IF(ISNUMBER(k_Recovery),"✓ Recovery score "&k_Recovery&" / 100.","· Recovery score appears once you log sleep, energy and stress."))',
        '=IF(k_Cardio7>4,"⚠ "&k_Cardio7&" cardio sessions this week → too much cardio eats into recovery and muscle. Keep it to 2–3.","✓ Cardio "&k_Cardio7&" / 2–3 this week.")',
        '=IF(AND(ISNUMBER(k_Energy7),k_Energy7<=4),"⚠ Energy low all week → treat it as a recovery week. If it keeps happening, talk to a doctor.",IF(ISNUMBER(k_Energy7),"✓ Energy "&TEXT(k_Energy7,"0.0")&" / 10 this week.","· Log energy 1–10 to track fatigue."))',
        '=IF(k_L1low>=3,"⚠ "&k_L1low&" non-negotiables slipped this week → SIMPLIFY THIS WEEK: run Minimum Viable Days, skip the optional list, add nothing new.","✓ No need to simplify. Keep the system the same size.")',
    ]
    for i, f in enumerate(rules):
        r = 23 + i
        put(ws, f"B{r}:L{r}", f, size=8, color=MUTED, bg=PANEL if i % 2 else PANEL2, indent=1)
    cf(ws, "B23:L36", 'LEFT($B23,1)="⚠"', color=YELLOW, fill=mix(YELLOW, PANEL, .1), stop=True)
    cf(ws, "B23:L36", 'LEFT($B23,2)="🤒"', color=PINK, fill=mix(PINK, PANEL, .15), stop=True)
    cf(ws, "B23:L36", 'LEFT($B23,1)="✓"', color=mix(GREEN, MUTED, .55), stop=True)

    # SYSTEM CHECK
    put(ws, "N22:S22", "SYSTEM CHECK", size=9, bold=True, color=TEAL, bg=PANEL, indent=1)
    diag = ('=IF(NOT(ISNUMBER(k_T_r7)),"Start logging",IF(k_L1low>=3,"The system is too complicated",'
            'IF(k_T_r7<Band3,"Reset week, not a failure",IF(k_T_r7<Band2,"Missed a few habits",'
            'IF(AND(ISNUMBER(k_Y_score),k_Y_score<0.5),"One off day","System working")))))')
    put(ws, "N23:S25", diag, size=13, bold=True, color=TEXT, bg=PANEL2, h="center", wrap=True)
    expl = ('=IF(NOT(ISNUMBER(k_T_r7)),"Tick today\'s habits on the month tab. The scores appear from the first logged day.",'
            'IF(k_L1low>=3,"3+ non-negotiables under 50% this week. That\'s a design problem, not a character problem. '
            'Run Minimum Viable Days for 7 days and drop one high-value habit.",'
            'IF(k_T_r7<Band3,"Rolling 7-day under 50%. Don\'t add motivation tactics. Shrink to the Minimum Viable Day '
            'until it\'s back above 70%.",IF(k_T_r7<Band2,"You missed some habits. That\'s data, not failure. Pick the one '
            'that matters most and protect it tomorrow.",IF(AND(ISNUMBER(k_Y_score),k_Y_score<0.5),"Yesterday was low, the '
            'rolling average is still healthy. Restart today. No make-up work, no punishment.","Consistency > intensity. '
            'Keep the system the same size; don\'t add habits just because it\'s going well.")))))')
    put(ws, "N26:S29", expl, size=7, color=MUTED, bg=PANEL2, wrap=True, v="top", indent=1)
    cf(ws, "N23:S25", 'OR($N$23="The system is too complicated",$N$23="Reset week, not a failure")', color=PINK, stop=True)
    cf(ws, "N23:S25", 'OR($N$23="Missed a few habits",$N$23="One off day")', color=YELLOW, stop=True)
    cf(ws, "N23:S25", '$N$23="System working"', color=GREEN, stop=True)
    checks = [
        ("Non-negotiables · 7d", "=k_L1avg7", "0%"),
        ("High-value vs targets · 7d", "=k_L2avg7", "0%"),
        ("Non-negotiables under 50%", "=k_L1low", "0"),
        ("Recovery score · 7d", '=IF(ISNUMBER(k_Recovery),k_Recovery,"—")', "0"),
        ("Days logged · 7d", '=k_Logged7&" / 7"', "@"),
    ]
    for i, (lab, f, fmt) in enumerate(checks):
        r = 30 + i
        put(ws, f"N{r}:Q{r}", lab, size=7, color=MUTED, bg=PANEL, indent=1)
        put(ws, f"R{r}:S{r}", f, size=8, bold=True, bg=PANEL, h="right", fmt=fmt)
    band_cf(ws, "R30:S31", "R30")
    put(ws, "N35:S36", "MINIMUM VIABLE DAY  ·  water · 10-min walk · protein as close as possible · 5-min mobility · "
                       "prayer · 10 min with your fiancée · lights out on time", size=7, color=YELLOW,
        bg=mix(YELLOW, PANEL, .1), wrap=True, indent=1)

    # THIS WEEK / THIS MONTH
    put(ws, "B38:J38", "THIS WEEK · last 7 days", size=9, bold=True, color=TEXT, bg=PANEL, indent=1)
    put(ws, "K38:S38", "THIS MONTH · last 30 days", size=9, bold=True, color=TEXT, bg=PANEL, indent=1)
    week = [
        ("Weight trend (7-day avg vs last week)", '=IF(ISNUMBER(k_Rate),TEXT(k_Rate,"+0.00;-0.00")&" kg · "&k_RateFlag,"—")'),
        ("Strength sessions", '=k_Str7&" / 4"'),
        ("Progression rate (14 days)", '=IF(ISNUMBER(k_ProgRate14),TEXT(k_ProgRate14,"0%"),"—")'),
        ("Core · calves · cardio", '=k_Core7&"/3 · "&k_Calf7&"/3 · "&k_Cardio7&"/2"'),
        ("Protein average", '=IF(ISNUMBER(k_Prot7),TEXT(k_Prot7,"0")&" g / "&ProtTarget,"—")'),
        ("Steps average", '=IF(ISNUMBER(k_Steps7),TEXT(k_Steps7,"#,##0")&" / "&TEXT(k_T_stepT,"#,##0"),"—")'),
        ("Sleep average", '=IF(ISNUMBER(k_Sleep7),TEXT(k_Sleep7,"0.0")&" h","—")'),
        ("Couple time · son · church · meal prep", '=k_Date7&" · "&k_Son7&" · "&k_Church7&" · "&k_Prep7'),
    ]
    month = [
        ("Weight change (7-day avg)", '=IF(ISNUMBER(k_Change30),TEXT(k_Change30,"+0.0;-0.0")&" kg","—")'),
        ("Waist (latest)", '=IF(ISNUMBER(k_Waist),TEXT(k_Waist,"0.0")&" cm"&IF(ISNUMBER(k_WaistPrev),"  ("&TEXT(k_Waist-k_WaistPrev,"+0.0;-0.0")&")",""),"measure every 2 weeks")'),
        ("Progress photos", '=IF(k_Photo30>0,"✓ taken","○ due this month")'),
        ("Strength sessions", '=k_Str30&" (target ~17)"'),
        ("Reading / learning days", '=k_Read30&" / 30"'),
        ("Prayer days", '=k_Pray30&" / 30"'),
        ("Goal weight (if muscle kept)", '=TEXT(GoalWeight,"0.0")&" kg"'),
        ("Weeks to goal at this pace", '=IF(ISNUMBER(k_WeeksToGoal),TEXT(k_WeeksToGoal,"0")&" weeks","needs a downward trend")'),
    ]
    for i in range(8):
        r = 39 + i
        bg = PANEL if i % 2 else PANEL2
        put(ws, f"B{r}:F{r}", week[i][0], size=7, color=MUTED, bg=bg, indent=1)
        put(ws, f"G{r}:J{r}", week[i][1], size=8, bold=True, bg=bg, h="right")
        put(ws, f"K{r}:O{r}", month[i][0], size=7, color=MUTED, bg=bg, indent=1)
        put(ws, f"P{r}:S{r}", month[i][1], size=8, bold=True, bg=bg, h="right")
    cf(ws, "P39:S46", 'LEFT(P39,1)="✓"', color=GREEN, stop=True)
    cf(ws, "P39:S46", 'LEFT(P39,1)="○"', color=YELLOW, stop=True)

    # PILLARS
    put(ws, "B48:S48", "PILLARS · 30-day consistency (high-value habits are scored against their weekly targets)",
        size=9, bold=True, color=TEXT, bg=PANEL, indent=1)
    lib = "'Habit Library'!"
    for i, (p, col) in enumerate(PILLARS.items()):
        r = 49 + (i % 4)
        n1, n2 = (("B", "D"), ("K", "M"))[i // 4]
        b1, b2 = (("E", "H"), ("N", "Q"))[i // 4]
        v1, v2 = (("I", "J"), ("R", "S"))[i // 4]
        pv = f'IFERROR(AVERAGEIFS({lib}$M$6:$M$40,{lib}$D$6:$D$40,"{p}"),"")'
        put(ws, f"{n1}{r}:{n2}{r}", p.upper(), size=8, bold=True, color=col, bg=PANEL, indent=1)
        put(ws, f"{v1}{r}:{v2}{r}", f"={pv}", size=9, bold=True, bg=PANEL, h="right", fmt="0%")
        put(ws, f"{b1}{r}:{b2}{r}", f"={bar(f'{v1}{r}', 16)}", size=8, color=col, bg=PANEL)
    band_cf(ws, "I49:J52", "$I49")
    band_cf(ws, "R49:S52", "$R49")

    # CHARTS
    calc = wb["Calc"]
    put(ws, "B54:J54", "FOUNDATION SCORE · last 30 days", size=8, bold=True, color=YELLOW, bg=PANEL, indent=1)
    put(ws, "K54:S54", "WEIGHT · daily vs 7-day average · 60 days", size=8, bold=True, color=VIOLET, bg=PANEL, indent=1)
    put(ws, "B71:J71", "STEPS vs adaptive target · 14 days", size=8, bold=True, color=TEAL, bg=PANEL, indent=1)
    put(ws, "K71:S71", "SLEEP · ENERGY · STRESS · MOOD · 14 days", size=8, bold=True, color=PINK, bg=PANEL, indent=1)
    paint(ws, "B55:S69", PANEL)
    paint(ws, "B72:S86", PANEL)

    fc = BarChart()
    fc.type = "col"
    fc.grouping = "clustered"
    fc.add_data(Reference(calc, min_col=5, min_row=34, max_row=63), titles_from_data=False)
    fc.series[0].graphicalProperties.solidFill = YELLOW
    fc.series[0].graphicalProperties.line.noFill = True
    fc.series[0].tx = None
    l7 = LineChart()
    l7.add_data(Reference(calc, min_col=6, min_row=34, max_row=63), titles_from_data=False)
    line_series(l7, TEAL, width=28000)
    fc.set_categories(Reference(calc, min_col=4, min_row=34, max_row=63))
    fc.gapWidth = 40
    style_chart(fc, ymin=0, ymax=1, yfmt="0%", major=0.25, legend=False)
    fc += l7
    size_chart(fc, ws, 2, 10, 55, 69)
    ws.add_chart(fc)
    chart_style("Dashboard", fc, "combo", [YELLOW, TEAL], vmin=0, vmax=1, fmt="percent")

    wc = LineChart()
    wc.add_data(Reference(calc, min_col=7, min_row=3, max_row=63), titles_from_data=True)
    wc.series[-1].graphicalProperties.line.noFill = True
    wc.series[-1].marker.symbol = "circle"
    wc.series[-1].marker.size = 4
    wc.series[-1].marker.graphicalProperties = GraphicalProperties(solidFill=VIOLET, ln=LineProperties(noFill=True))
    wc.add_data(Reference(calc, min_col=8, min_row=3, max_row=63), titles_from_data=True)
    line_series(wc, TEAL, width=28000)
    wc.set_categories(Reference(calc, min_col=4, min_row=4, max_row=63))
    style_chart(wc, yfmt="0.0")
    size_chart(wc, ws, 11, 19, 55, 69)
    ws.add_chart(wc)
    chart_style("Dashboard", wc, "weight", [VIOLET, TEAL])

    sc = BarChart()
    sc.type = "col"
    s = Series(Reference(calc, min_col=9, min_row=50, max_row=63), title="Steps")
    s.graphicalProperties.solidFill = TEAL
    s.graphicalProperties.line.noFill = True
    sc.series.append(s)
    st = LineChart()
    s2 = Series(Reference(calc, min_col=10, min_row=50, max_row=63), title="Target")
    st.series.append(s2)
    line_series(st, YELLOW, width=19000, smooth=False, dash="dash")
    sc.set_categories(Reference(calc, min_col=4, min_row=50, max_row=63))
    sc.gapWidth = 50
    style_chart(sc, ymin=0, yfmt="#,##0")
    sc += st
    size_chart(sc, ws, 2, 10, 72, 86)
    ws.add_chart(sc)
    chart_style("Dashboard", sc, "combo", [TEAL, YELLOW])

    mc = LineChart()
    for col_i, title, col in ((11, "Sleep h", BLUE), (12, "Energy", YELLOW), (13, "Stress", PINK), (14, "Mood", TEAL)):
        mc.series.append(Series(Reference(calc, min_col=col_i, min_row=50, max_row=63), title=title))
        line_series(mc, col)
    mc.set_categories(Reference(calc, min_col=4, min_row=50, max_row=63))
    style_chart(mc, ymin=0, ymax=10, major=2)
    size_chart(mc, ws, 11, 19, 72, 86)
    ws.add_chart(mc)
    chart_style("Dashboard", mc, "line", [BLUE, YELLOW, PINK, TEAL], vmin=0, vmax=10)
    ws.freeze_panes = "A5"


# ======================================================================== WEEK PLAN
def build_week(wb, ws):
    sn = ws.title
    page(ws)
    ws.sheet_properties.tabColor = BLUE
    paint(ws, "A1:V66", BG)
    ws.column_dimensions["A"].width = 2
    base = []
    for d in range(7):
        c = 2 + d * 3
        base.append(c)
        ws.column_dimensions[CL(c)].width = 15
        ws.column_dimensions[CL(c + 1)].width = 4
        ws.column_dimensions[CL(c + 2)].width = 1.6
    heights(ws, range(1, 67), 17)
    put(ws, "B2:H2", "WEEK PLAN", size=24, bold=True)
    put(ws, "B3:K3", "Plan the week · tick tasks here · check-ins and scores pull from your daily log", size=8, color=MUTED)
    put(ws, "N2:O2", "WEEK STARTING", size=7, bold=True, color=MUTED, bg=PANEL, h="right")
    put(ws, "Q2:R2", f"={TODAY}-WEEKDAY({TODAY},3)", size=10, bold=True, color=YELLOW, bg=PANEL2, h="center",
        fmt="ddd d mmm yyyy")
    name(wb, "WeekStart", f"'{sn}'!$Q$2")
    put(ws, "N3:R3", "type a Monday to plan ahead", size=6, color=DIM, h="right")
    ws.row_dimensions[2].height = 26

    # helper table
    hr0 = 56
    put(ws, f"B{hr0 - 1}:N{hr0 - 1}", "chart data below (kept invisible on purpose)", size=6, color=mix(DIM, BG, .5))

    def val(key, dexpr):
        return f'IFERROR(INDEX({DR(key)},MATCH({dexpr},{DR("date")},0)),"")'

    for d in range(7):
        r = hr0 + 1 + d
        dexpr = f"(WeekStart+{d})"
        ws[f"B{r}"] = f'=TEXT({dexpr},"ddd")'
        ws[f"C{r}"] = f"={dexpr}"
        ws[f"C{r}"].number_format = "d mmm"
        ws[f"D{r}"] = "=" + val("score", dexpr)
        ws[f"D{r}"].number_format = "0%"
        ws[f"E{r}"] = f"=IF(ISNUMBER(D{r}),D{r},NA())"
        ws[f"F{r}"] = f"=IF(ISNUMBER(D{r}),D{r},0)"
        ws[f"G{r}"] = f"=1-F{r}"
        for j, key in enumerate(("energy", "stress", "mood")):
            ws[f"{CL(8 + j)}{r}"] = f'=IF(ISNUMBER({val(key, dexpr)}),{val(key, dexpr)},NA())'
        for col in "BCDEFGHIJ":
            ws[f"{col}{r}"].font = F(6, BG)
    r = hr0 + 8
    ws[f"B{r}"] = "week"
    ws[f"D{r}"] = f'=IFERROR(AVERAGE(D{hr0 + 1}:D{hr0 + 7}),"")'
    ws[f"F{r}"] = f"=IF(ISNUMBER(D{r}),D{r},0)"
    ws[f"G{r}"] = f"=1-F{r}"
    for col in "BDFG":
        ws[f"{col}{r}"].font = F(6, BG)

    # top charts
    put(ws, "B5:H5", "DAILY FOUNDATION SCORE", size=8, bold=True, color=TEXT, bg=PANEL, indent=1)
    put(ws, "K5:M5", "WEEK", size=8, bold=True, color=TEXT, bg=PANEL, h="center")
    put(ws, "N5:U5", "CHECK-IN TREND", size=8, bold=True, color=TEXT, bg=PANEL, indent=1)
    paint(ws, "B6:U15", PANEL)
    bc = BarChart()
    bc.type = "col"
    s = Series(Reference(ws, min_col=5, min_row=hr0 + 1, max_row=hr0 + 7), title="Score")
    bc.series.append(s)
    for i in range(7):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = DAY_COLORS[i]
        pt.graphicalProperties.line.noFill = True
        s.dPt.append(pt)
    bc.set_categories(Reference(ws, min_col=2, min_row=hr0 + 1, max_row=hr0 + 7))
    bc.gapWidth = 45
    style_chart(bc, ymin=0, ymax=1, yfmt="0%", major=0.25, legend=False)
    size_chart(bc, ws, 2, 9, 6, 15)
    ws.add_chart(bc)
    chart_style(sn, bc, "bar", [], vmin=0, vmax=1, fmt="percent")
    dn = DoughnutChart()
    dn.holeSize = 70
    s = Series(Reference(ws, min_col=6, max_col=7, min_row=hr0 + 8, max_row=hr0 + 8), title="Week")
    dn.series.append(s)
    for i, col in enumerate((TEAL, PANEL2)):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = col
        pt.graphicalProperties.line.solidFill = PANEL
        s.dPt.append(pt)
    dn.title = None
    dn.legend = None
    dn.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
    dn.plot_area.graphicalProperties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
    size_chart(dn, ws, 11, 13, 6, 13)
    ws.add_chart(dn)
    chart_style(sn, dn, "donut", [TEAL, PANEL2])
    put(ws, "K14:M15", f'=IF(ISNUMBER(D{hr0 + 8}),TEXT(D{hr0 + 8},"0%"),"—")', size=16, bold=True, color=TEAL,
        bg=PANEL, h="center")
    lc = LineChart()
    for j, (title, col) in enumerate((("Energy", YELLOW), ("Stress", PINK), ("Mood", TEAL))):
        lc.series.append(Series(Reference(ws, min_col=8 + j, min_row=hr0 + 1, max_row=hr0 + 7), title=title))
        line_series(lc, col)
    lc.set_categories(Reference(ws, min_col=2, min_row=hr0 + 1, max_row=hr0 + 7))
    style_chart(lc, ymin=0, ymax=10, major=2)
    size_chart(lc, ws, 14, 21, 6, 15)
    ws.add_chart(lc)
    chart_style(sn, lc, "line", [YELLOW, PINK, TEAL], vmin=0, vmax=10)

    # day panels
    for d in range(7):
        c1, c2 = CL(base[d]), CL(base[d] + 1)
        col = DAY_COLORS[d]
        dexpr = f"(WeekStart+{d})"
        put(ws, f"{c1}17:{c2}17", f'=UPPER(TEXT({dexpr},"dddd"))', size=10, bold=True, color=BG, bg=col, h="center")
        put(ws, f"{c1}18:{c2}18", f"={dexpr}", size=8, color=TEXT, bg=mix(col, PANEL, .3), h="center", fmt="d mmm yyyy")
        put(ws, f"{c1}19:{c2}20", f'="06:30 · "&INDEX(PlanTraining,{d + 1})', size=7, color=col, bg=PANEL, h="center",
            wrap=True)
        paint(ws, f"{c1}21:{c2}27", PANEL)
        dc = DoughnutChart()
        dc.holeSize = 72
        s = Series(Reference(ws, min_col=6, max_col=7, min_row=hr0 + 1 + d, max_row=hr0 + 1 + d), title="Score")
        dc.series.append(s)
        for i, pc in enumerate((col, PANEL2)):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = pc
            pt.graphicalProperties.line.solidFill = PANEL
            s.dPt.append(pt)
        dc.legend = None
        dc.title = None
        dc.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
        dc.plot_area.graphicalProperties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
        size_chart(dc, ws, base[d], base[d] + 1, 21, 27)
        ws.add_chart(dc)
        chart_style(sn, dc, "donut", [col, PANEL2])
        put(ws, f"{c1}28:{c2}28", f'=IF(ISNUMBER(D{hr0 + 1 + d}),TEXT(D{hr0 + 1 + d},"0%"),"—")', size=14, bold=True,
            color=col, bg=PANEL, h="center")
        put(ws, f"{c1}29:{c2}29", "foundation score", size=6, color=MUTED, bg=PANEL, h="center", v="top")
        put(ws, f"{c1}30:{c2}30", "TOP 3 PRIORITIES", size=7, bold=True, color=col, bg=mix(col, PANEL, .18), indent=1)
        for r in range(31, 34):
            put(ws, f"{c1}{r}", size=7, bg=PANEL2, indent=1, border=TILE)
            put(ws, f"{c2}{r}", size=9, color=col, bg=PANEL2, h="center", border=TILE)
        put(ws, f"{c1}34:{c2}34", "TASKS", size=7, bold=True, color=col, bg=mix(col, PANEL, .18), indent=1)
        for r in range(35, 43):
            put(ws, f"{c1}{r}", size=7, bg=PANEL, indent=1, border=TILE)
            put(ws, f"{c2}{r}", size=9, color=col, bg=PANEL, h="center", border=TILE)
        checkbox(sn, f"{c2}31:{c2}33")
        checkbox(sn, f"{c2}35:{c2}42")
        done = f"COUNTIF({c2}31:{c2}42,TRUE)+COUNTIF({c2}31:{c2}42,\"x\")"
        put(ws, f"{c1}43", "Completed", size=7, color=MUTED, bg=PANEL, indent=1)
        put(ws, f"{c2}43", f"={done}", size=8, bold=True, color=GREEN, bg=PANEL, h="center")
        put(ws, f"{c1}44", "Not completed", size=7, color=MUTED, bg=PANEL, indent=1)
        put(ws, f"{c2}44", f"=MAX(0,COUNTA({c1}31:{c1}33,{c1}35:{c1}42)-({done}))", size=8, bold=True, color=PINK,
            bg=PANEL, h="center")
        put(ws, f"{c1}45:{c2}45", "CHECK-IN · from daily log", size=6, bold=True, color=col, bg=mix(col, PANEL, .18),
            indent=1)
        for i, (lab, key, mx, colr) in enumerate((("Sleep", "sleeph", 9, BLUE), ("Energy", "energy", 10, YELLOW),
                                                  ("Stress", "stress", 10, PINK), ("Mood", "mood", 10, TEAL),
                                                  ("Protein", "protein", 150, ORANGE), ("Steps", "steps", 10000, GREEN))):
            r = 46 + i
            v = val(key, dexpr)
            put(ws, f"{c1}{r}", f'="{lab:<8}"&IF(ISNUMBER({v}),{bar(f"({v})/{mx}", 6)},"")', size=7, color=colr,
                bg=PANEL, indent=1)
            fmt = {"sleeph": "0.0", "steps": '0.0"k"', "protein": "0"}.get(key, "0")
            shown = f"({v})/1000" if key == "steps" else v
            put(ws, f"{c2}{r}", f'=IF(ISNUMBER({v}),{shown},"")', size=7, bold=True, bg=PANEL, h="center", fmt=fmt)
        cf(ws, f"{c1}31:{c1}42", f'OR({c2}31=TRUE,{c2}31="x")', color=DIM)
    ws.freeze_panes = "A4"


# ======================================================================== BODY
def build_body(wb, ws):
    page(ws)
    ws.sheet_properties.tabColor = PINK
    paint(ws, "A1:W48", BG)
    widths(ws, {"A": 2, "B": 13, "C": 10, "D": 9, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 9, "K": 26, "L": 2})
    for c in range(13, 24):
        ws.column_dimensions[CL(c)].width = 8
    heights(ws, range(1, 49), 18)
    put(ws, "B2:H2", "BODY", size=24, bold=True)
    put(ws, "B3:K3", "Fat down, muscle kept. Weigh daily, decide weekly, measure every two weeks, photos monthly.",
        size=8, color=MUTED)
    ws.row_dimensions[2].height = 30
    prof = [
        ("Age · height", '=Age&" · "&HeightCm&" cm"'),
        ("Starting weight", '=TEXT(StartWeight,"0.0")&" kg"'),
        ("Body fat (est.)", '=TEXT(StartBF,"0%")&" → goal "&TEXT(GoalBF,"0%")'),
        ("Lean mass (est.)", '=TEXT(LeanMass,"0.0")&" kg"'),
        ("Goal weight if muscle is kept", '=TEXT(GoalWeight,"0.0")&" kg"'),
        ("Current 7-day average", '=IF(ISNUMBER(k_T_w7),TEXT(k_T_w7,"0.0")&" kg","—")'),
        ("Healthy pace", '="−"&TEXT(LossMin,"0.0")&" to −"&TEXT(LossMax,"0.0")&" kg / week"'),
        ("Time at that pace", '=ROUND((StartWeight-GoalWeight)/LossMax,0)&"–"&ROUND((StartWeight-GoalWeight)/LossMin,0)&" weeks"'),
    ]
    put(ws, "B5:F5", "PROFILE & GOAL", size=8, bold=True, color=PINK, bg=mix(PINK, PANEL, .16), indent=1)
    for i, (lab, f) in enumerate(prof):
        r = 6 + i
        put(ws, f"B{r}:D{r}", lab, size=8, color=MUTED, bg=PANEL, indent=1)
        put(ws, f"E{r}:F{r}", f, size=8, bold=True, bg=PANEL, h="right")
    put(ws, "H5:K5", "PROGRESS TO GOAL WEIGHT", size=8, bold=True, color=TEAL, bg=mix(TEAL, PANEL, .16), indent=1)
    prog = '(StartWeight-k_T_w7)/(StartWeight-GoalWeight)'
    put(ws, "H6:K7", f'=IF(ISNUMBER(k_T_w7),TEXT(MAX(0,{prog}),"0%"),"—")', size=20, bold=True, color=TEAL, bg=PANEL,
        h="center")
    put(ws, "H8:K8", f'=IF(ISNUMBER(k_T_w7),{bar(f"MAX(0,{prog})", 20)},"")', size=8, color=TEAL, bg=PANEL, h="center")
    put(ws, "H9:K13", "Estimates only, assuming lean mass is kept. The real test is the waist, the mirror and "
                      "strength going up. If the scale stalls but the waist drops, it's working.", size=7,
        color=MUTED, bg=PANEL, wrap=True, v="top", indent=1)

    hdr = ["Date", "Weight 7d", "Waist", "Chest", "Arms", "Thigh", "Calf", "Neck", "Photos", "Notes"]
    put(ws, "B15:K15", "MEASUREMENTS · every 2 weeks (cm) · photos monthly: same light, distance, pose, time · front / side / back",
        size=8, bold=True, color=TEXT, bg=PANEL, indent=1)
    for i, hd in enumerate(hdr):
        put(ws, f"{CL(2 + i)}16", hd, size=7, bold=True, color=MUTED, bg=PANEL2, h="center")
    for i in range(27):
        r = 17 + i
        bg = PANEL if i % 2 else PANEL2
        put(ws, f"B{r}", f"=TrackStart+{14 * i}", size=8, bg=bg, h="center", fmt="d mmm yy")
        put(ws, f"C{r}", f'=IFERROR(INDEX({DR("w7")},MATCH(B{r},{DR("date")},0)),"")', size=8, color=MUTED, bg=bg,
            h="center", fmt="0.0")
        for col in "DEFGHI":
            put(ws, f"{col}{r}", size=8, color=YELLOW, bg=bg, h="center", fmt="0.0")
        put(ws, f"J{r}", size=9, color=PINK, bg=bg, h="center")
        put(ws, f"K{r}", size=7, color=MUTED, bg=bg, indent=1)
    checkbox("Body", "J17:J43")
    put(ws, "B45:K46", "Dates are suggestions (every 14 days from your start date). Overwrite any date with the day you "
                       "actually measured. Waist is the one that matters most.", size=7, color=MUTED, wrap=True,
        italic=True)

    calc = wb["Calc"]
    put(ws, "M5:W5", "WEIGHT · weekly average", size=8, bold=True, color=VIOLET, bg=PANEL, indent=1)
    paint(ws, "M6:W24", PANEL)
    lc = LineChart()
    lc.add_data(Reference(calc, min_col=17, min_row=3, max_row=56), titles_from_data=True)
    line_series(lc, VIOLET, width=28000)
    lc.series[-1].marker.symbol = "circle"
    lc.series[-1].marker.size = 5
    lc.series[-1].marker.graphicalProperties = GraphicalProperties(solidFill=VIOLET, ln=LineProperties(noFill=True))
    lc.set_categories(Reference(calc, min_col=16, min_row=4, max_row=56))
    style_chart(lc, yfmt="0.0", legend=False)
    size_chart(lc, ws, 13, 23, 6, 24)
    ws.add_chart(lc)
    chart_style("Body", lc, "line", [VIOLET])
    put(ws, "M26:W26", "WAIST · cm", size=8, bold=True, color=TEAL, bg=PANEL, indent=1)
    paint(ws, "M27:W44", PANEL)
    wc = LineChart()
    wc.add_data(Reference(calc, min_col=20, min_row=3, max_row=30), titles_from_data=True)
    line_series(wc, TEAL, width=28000)
    wc.series[-1].marker.symbol = "circle"
    wc.series[-1].marker.size = 5
    wc.series[-1].marker.graphicalProperties = GraphicalProperties(solidFill=TEAL, ln=LineProperties(noFill=True))
    wc.set_categories(Reference(calc, min_col=19, min_row=4, max_row=30))
    style_chart(wc, yfmt="0.0", legend=False)
    size_chart(wc, ws, 13, 23, 27, 44)
    ws.add_chart(wc)
    chart_style("Body", wc, "line", [TEAL])
    ws.freeze_panes = "A4"


# ======================================================================== WEEKLY REVIEW
def build_weekly(wb, ws):
    page(ws)
    ws.sheet_properties.tabColor = GREEN
    n = 53
    last = 8 + n - 1
    paint(ws, f"A1:AK{last + 3}", BG)
    auto = [("Week", 6, "0"), ("Monday", 9, "d mmm"), ("Foundation", 9, "0%"), ("Status", 11, "@"),
            ("Avg kg", 7, "0.0"), ("Δ kg", 7, "+0.0;-0.0"), ("Pace", 9, "@"), ("Strength /4", 7, "0"),
            ("Progress %", 7, "0%"), ("Core /3", 6, "0"), ("Calves /3", 6, "0"), ("Cardio /2", 6, "0"),
            ("Steps avg", 8, "#,##0"), ("Protein avg", 7, "0"), ("Sleep avg", 7, "0.0"), ("Energy", 6, "0.0"),
            ("Stress", 6, "0.0"), ("Prayer /7", 6, "0"), ("Church", 6, "0"), ("Couple", 6, "0"), ("Son", 6, "0"),
            ("Family", 6, "0"), ("Meal prep", 6, "0"), ("Deep work", 6, "0"), ("Recovery", 7, "0")]
    mine = [("Business ✓", 7), ("Finance ✓", 7), ("Major wins", 24), ("Major problems", 24),
            ("What I learned", 22), ("Spirit / relationships", 24), ("ONE change for next week", 26),
            ("Next week's priority", 22), ("Done ✓", 6)]
    ws.column_dimensions["A"].width = 2
    col = 2
    cols = {}
    for nm, w, _ in auto:
        ws.column_dimensions[CL(col)].width = w
        cols[nm] = CL(col)
        col += 1
    for nm, w in mine:
        ws.column_dimensions[CL(col)].width = w
        cols[nm] = CL(col)
        col += 1
    lastcol = CL(col - 1)
    heights(ws, range(1, last + 4), 20)
    put(ws, "B2", "WEEKLY REVIEW", size=24, bold=True)
    ws.row_dimensions[2].height = 30
    put(ws, "B3", "Sunday evening, 15–30 minutes. The left side fills itself from your daily log; you only write "
                     "the right side. End with ONE change — only one.", size=8, color=MUTED)
    paint(ws, "B4:Z4", mix(YELLOW, PANEL, .1))
    put(ws, "B4", "Business check (tick Business ✓ when done): revenue · pipeline · sales activity · client "
                     "delivery · marketing · cash flow · outstanding tasks · strategic priority.   Life: laundry · "
                     "rooms · bathroom · groceries · meal prep · finances · calendar · desk ergonomics check.",
        size=7, color=YELLOW, bg=mix(YELLOW, PANEL, .1), indent=1)
    a_first, a_last = cols["Week"], cols["Recovery"]
    m_first = cols["Business ✓"]
    paint(ws, f"B6:C6", PANEL)
    put(ws, f"D6:{a_last}6", "AUTO · from your daily log", size=7, bold=True, color=TEAL,
        bg=mix(TEAL, PANEL, .16), h="center")
    put(ws, f"{m_first}6:{lastcol}6", "YOUR REVIEW", size=7, bold=True, color=YELLOW, bg=mix(YELLOW, PANEL, .16),
        h="center")
    for nm in list(cols):
        put(ws, f"{cols[nm]}7", nm, size=6, bold=True, color=MUTED, bg=PANEL2, h="center", wrap=True)
    ws.row_dimensions[7].height = 26

    def sumw(key, r):
        return f'SUMIFS({DR(key)},{DR("date")},">="&$C{r},{DR("date")},"<="&($C{r}+6))'

    def avgw(key, r):
        return f'IFERROR(AVERAGEIFS({DR(key)},{DR("date")},">="&$C{r},{DR("date")},"<="&($C{r}+6),{DR(key)},">0"),"")'

    for i in range(n):
        r = 8 + i
        bg = PANEL if i % 2 else PANEL2
        f = {
            "Week": i + 1,
            "Monday": f"=DATE({FIRST_MONDAY.year},{FIRST_MONDAY.month},{FIRST_MONDAY.day})+{7 * i}",
            "Foundation": f'=IFERROR(AVERAGEIFS({DR("score")},{DR("date")},">="&$C{r},{DR("date")},"<="&($C{r}+6)),"")',
        }
        f["Status"] = f"={status(cols['Foundation'] + str(r))}"
        f["Avg kg"] = "=" + avgw("weight", r)
        f["Δ kg"] = (f'=IF(AND(ISNUMBER({cols["Avg kg"]}{r}),ISNUMBER({cols["Avg kg"]}{r - 1})),'
                     f'{cols["Avg kg"]}{r}-{cols["Avg kg"]}{r - 1},"")') if i else '=""'
        dk, ak = f'{cols["Δ kg"]}{r}', f'{cols["Avg kg"]}{r}'
        f["Pace"] = (f'=IF(NOT(ISNUMBER({dk})),"",IF(-{dk}>LossMaxPct*{ak},"too fast",IF(-{dk}>=LossMin,'
                     f'IF(-{dk}<=LossMax,"on target","fast end"),IF(-{dk}>0,"slow","flat / up"))))')
        f["Strength /4"] = "=" + sumw("str", r)
        f["Progress %"] = f'=IF({cols["Strength /4"]}{r}=0,"",MIN(1,{sumw("prog", r)}/{cols["Strength /4"]}{r}))'
        f["Core /3"] = "=" + sumw("core", r)
        f["Calves /3"] = "=" + sumw("calf", r)
        f["Cardio /2"] = f'={sumw("cardio", r)}+{sumw("xcardio", r)}'
        f["Steps avg"] = "=" + avgw("steps", r)
        f["Protein avg"] = "=" + avgw("protein", r)
        f["Sleep avg"] = "=" + avgw("sleeph", r)
        f["Energy"] = "=" + avgw("energy", r)
        f["Stress"] = "=" + avgw("stress", r)
        f["Prayer /7"] = "=" + sumw("pray", r)
        f["Church"] = "=" + sumw("church", r)
        f["Couple"] = "=" + sumw("couple", r)
        f["Son"] = "=" + sumw("son", r)
        f["Family"] = "=" + sumw("family", r)
        f["Meal prep"] = "=" + sumw("prep", r)
        f["Deep work"] = "=" + sumw("deep", r)
        sl, en, stc = (f'{cols[x]}{r}' for x in ("Sleep avg", "Energy", "Stress"))
        stv, pg = f'{cols["Strength /4"]}{r}', f'{cols["Progress %"]}{r}'
        f["Recovery"] = (f'=IF(NOT(ISNUMBER({cols["Foundation"]}{r})),"",ROUND(100*(0.3*IF(ISNUMBER({sl}),MIN(1,{sl}/SleepTarget),0.7)+'
                         f'0.2*IF(ISNUMBER({en}),{en}/10,0.7)+0.2*IF(ISNUMBER({stc}),(10-{stc})/9,0.7)+'
                         f'0.15*IF({stv}<7,1,0)+0.15*IF(ISNUMBER({pg}),MIN(1,{pg}/0.5),0.5)),0))')
        for nm, w, fmt in auto:
            v = f[nm]
            if nm not in ("Week", "Monday") and isinstance(v, str) and v.startswith("=") and v != '=""':
                v = f'=IF($C{r}>$AN$1,"",{v[1:]})'
            put(ws, f"{cols[nm]}{r}", v, size=7, bold=nm in ("Foundation", "Status"), bg=bg, h="center", fmt=fmt)
        for nm, w in mine:
            put(ws, f"{cols[nm]}{r}", size=7, color=YELLOW if "✓" not in nm else GREEN, bg=bg,
                h="center" if "✓" in nm else "left", wrap=True, indent=0 if "✓" in nm else 1)
    for nm in ("Business ✓", "Finance ✓", "Done ✓"):
        checkbox("Weekly Review", f"{cols[nm]}8:{cols[nm]}{last}")
    band_cf(ws, f"{cols['Foundation']}8:{cols['Foundation']}{last}", f"{cols['Foundation']}8")
    status_cf(ws, f"{cols['Status']}8:{cols['Status']}{last}", f"{cols['Status']}8")
    pc = cols["Pace"]
    cf(ws, f"{pc}8:{pc}{last}", f'{pc}8="on target"', color=GREEN, stop=True)
    cf(ws, f"{pc}8:{pc}{last}", f'OR({pc}8="too fast",{pc}8="flat / up")', color=PINK, stop=True)
    cf(ws, f"{pc}8:{pc}{last}", f'OR({pc}8="slow",{pc}8="fast end")', color=YELLOW, stop=True)
    cf(ws, f"B8:{lastcol}{last}", f'AND($C8<=$AN$1,$C8+6>=$AN$1)', fill=mix(YELLOW, PANEL, .18))
    ws["AN1"] = "=AsOf"
    ws["AN1"].font = F(6, BG)
    name(wb, "AsOfW", "'Weekly Review'!$AN$1")
    ws.freeze_panes = "D8"


# ======================================================================== MONTHLY REVIEW
def build_monthly(wb, ws):
    page(ws)
    ws.sheet_properties.tabColor = ORANGE
    paint(ws, "A1:AN24", BG)
    auto = [("Month", 10, "mmm yyyy"), ("Foundation", 9, "0%"), ("Status", 11, "@"), ("Weight start", 8, "0.0"),
            ("Weight end", 8, "0.0"), ("Change", 7, "+0.0;-0.0"), ("Waist", 7, "0.0"), ("Strength", 7, "0"),
            ("Progress %", 7, "0%"), ("Sleep avg", 7, "0.0"), ("Energy", 7, "0.0"), ("Steps avg", 8, "#,##0"),
            ("Protein avg", 7, "0"), ("Prayer days", 7, "0"), ("Church", 6, "0"), ("Couple", 6, "0"), ("Son", 6, "0"),
            ("Read days", 6, "0")]
    mine = [("Photos ✓", 7), ("Body fat est.", 8), ("Books done", 16), ("Skills learned", 18), ("Revenue", 10),
            ("Profit", 10), ("Expenses", 10), ("New clients", 8), ("Retention", 9), ("Major wins", 22),
            ("Major failures", 22), ("What worked", 20), ("What didn't", 20), ("Stop", 16), ("Start", 16),
            ("Continue", 16), ("More disciplined, loving, truthful, patient, useful?", 28),
            ("Next month's focus", 22)]
    ws.column_dimensions["A"].width = 2
    col, cols = 2, {}
    for nm, w, _ in auto:
        ws.column_dimensions[CL(col)].width = w
        cols[nm] = CL(col)
        col += 1
    for nm, w in mine:
        ws.column_dimensions[CL(col)].width = w
        cols[nm] = CL(col)
        col += 1
    lastcol = CL(col - 1)
    heights(ws, range(1, 25), 30)
    put(ws, "B2", "MONTHLY REVIEW", size=24, bold=True)
    put(ws, "B3", "Once a month. Numbers fill themselves; you write the reflection. Stop · start · continue, then "
                     "one focus for next month.", size=8, color=MUTED)
    ws.row_dimensions[3].height = 18
    paint(ws, "B6:B6", PANEL)
    put(ws, f"C6:{cols['Read days']}6", "AUTO · from your daily log and Body tab", size=7, bold=True,
        color=TEAL, bg=mix(TEAL, PANEL, .16), h="center")
    put(ws, f"{cols['Photos ✓']}6:{lastcol}6", "YOUR REVIEW", size=7, bold=True, color=YELLOW,
        bg=mix(YELLOW, PANEL, .16), h="center")
    ws.row_dimensions[6].height = 18
    for nm in cols:
        put(ws, f"{cols[nm]}7", nm, size=6, bold=True, color=MUTED, bg=PANEL2, h="center", wrap=True)
    ws.row_dimensions[7].height = 34
    for k, y, m, sn in MONTHS:
        r = 8 + k
        bg = PANEL if k % 2 else PANEL2
        start = f"DATE({y},{m},1)"
        end = f"EOMONTH({start},0)"

        def sumk(key):
            return f'SUMIFS({DR(key)},{DR("month")},{k + 1})'

        def avgk(key):
            return f'IFERROR(AVERAGEIFS({DR(key)},{DR("month")},{k + 1},{DR(key)},">0"),"")'

        wd = f'_xlfn.MAXIFS(Body!$B$17:$B$43,Body!$B$17:$B$43,">="&{start},Body!$B$17:$B$43,"<="&{end},Body!$D$17:$D$43,">0")'
        f = {
            "Month": f"={start}",
            "Foundation": f'=IFERROR(AVERAGEIFS({DR("score")},{DR("month")},{k + 1}),"")',
            "Status": "=" + status(f"{cols['Foundation']}{r}"),
            "Weight start": f'=IFERROR(AVERAGEIFS({DR("weight")},{DR("date")},">="&{start},{DR("date")},"<="&({start}+6),{DR("weight")},">0"),"")',
            "Weight end": f'=IFERROR(AVERAGEIFS({DR("weight")},{DR("date")},">="&({end}-6),{DR("date")},"<="&{end},{DR("weight")},">0"),"")',
            "Change": f'=IF(AND(ISNUMBER({cols["Weight start"]}{r}),ISNUMBER({cols["Weight end"]}{r})),{cols["Weight end"]}{r}-{cols["Weight start"]}{r},"")',
            "Waist": f'=IF({wd}=0,"",INDEX(Body!$D$17:$D$43,MATCH({wd},Body!$B$17:$B$43,0)))',
            "Strength": "=" + sumk("str"),
            "Progress %": f'=IF({cols["Strength"]}{r}=0,"",MIN(1,{sumk("prog")}/{cols["Strength"]}{r}))',
            "Sleep avg": "=" + avgk("sleeph"),
            "Energy": "=" + avgk("energy"),
            "Steps avg": "=" + avgk("steps"),
            "Protein avg": "=" + avgk("protein"),
            "Prayer days": "=" + sumk("pray"),
            "Church": "=" + sumk("church"),
            "Couple": "=" + sumk("couple"),
            "Son": "=" + sumk("son"),
            "Read days": "=" + sumk("c_read"),
        }
        for nm, w, fmt in auto:
            v = f[nm]
            if nm != "Month":
                v = f'=IF({start}>AsOf,"",{v[1:]})'
            put(ws, f"{cols[nm]}{r}", v, size=7, bold=nm in ("Month", "Foundation", "Status"), bg=bg, h="center",
                fmt=fmt)
        for nm, w in mine:
            fmt = "#,##0" if nm in ("Revenue", "Profit", "Expenses") else ("0%" if nm == "Body fat est." else None)
            put(ws, f"{cols[nm]}{r}", size=7, color=GREEN if "✓" in nm else YELLOW, bg=bg,
                h="center" if ("✓" in nm or fmt) else "left", wrap=True, fmt=fmt, indent=0 if ("✓" in nm or fmt) else 1)
    checkbox("Monthly Review", f"{cols['Photos ✓']}8:{cols['Photos ✓']}19")
    band_cf(ws, f"{cols['Foundation']}8:{cols['Foundation']}19", f"{cols['Foundation']}8")
    status_cf(ws, f"{cols['Status']}8:{cols['Status']}19", f"{cols['Status']}8")
    put(ws, "C21:Z22", "Faith reflection: the tracker measures practice, not faith. Use the question column honestly, "
                       "then pick one next step. Business numbers: revenue, profit, expenses, new clients, retention.",
        size=7, color=MUTED, italic=True, wrap=True)
    ws.freeze_panes = "C8"


# ======================================================================== TEXT SHEETS
ROUTINES = [
    ("MORNING RESET · 06:00–06:20 · one tick", TEAL, [
        ("06:00", "Out of bed straight away · drink water · make the bed · open curtains · bathroom · weigh in · no social media"),
        ("Hydration", "500–750 ml to start. Don't force huge amounts."),
        ("Light", "5–15 minutes outside, not through a window."),
        ("Prayer", "5–10 min: prayer · Scripture · one thing I'm grateful for · one thing I need God's help with today."),
    ]),
    ("MOBILITY & POSTURE · 06:15–06:25 · one tick", BLUE, [
        ("Neck", "Chin tucks × 10"),
        ("Upper back", "Wall angels × 10 · thoracic extensions × 8–10 over a chair back"),
        ("Shoulders", "External rotation with 2 kg × 12–15 · scapular retractions × 12–15"),
        ("Chest & hips", "Doorway chest stretch 2 × 30 s · cat-cow × 8 · hip-flexor stretch 30 s / side"),
        ("Rule", "Consistency, not administrative gymnastics. One checkbox for the whole routine."),
    ]),
    ("TRAINING SPLIT · 06:30–07:30", PINK, [
        ("Monday", "Upper body + posture"),
        ("Tuesday", "Lower body + calves + core"),
        ("Wednesday", "Recovery: walk + mobility"),
        ("Thursday", "Upper body + posture"),
        ("Friday", "Lower body + calves + core"),
        ("Saturday", "Optional cardio / walk / mobility"),
        ("Sunday", "Recovery · church · meal prep · review"),
    ]),
    ("UPPER BODY + POSTURE · Mon / Thu", PINK, [
        ("Push-ups", "4 × 6–15 · progress incline → standard → decline → deficit"),
        ("Pike push-ups", "3 × 5–10"),
        ("One-arm row", "10 kg · 4 × 8–15 / side · pause at the top"),
        ("Floor press", "10 kg single-arm · 3 × 8–12 / side"),
        ("Rear delts", "Reverse fly 2 × 2 kg · 3 × 15–20 · Y-raise 2 × 12–15"),
        ("Scap work", "Scap push-ups 2 × 10–12 · finish with wall angels"),
    ]),
    ("LOWER BODY + CALVES + CORE · Tue / Fri", PINK, [
        ("Split squat", "Bulgarian split squat, 10 kg · 3–4 × 8–12 / leg"),
        ("Goblet squat", "10 kg · 3 × 12–20 · 3-second lowering"),
        ("Hinge", "Single-leg Romanian deadlift, 10 kg · 3 × 8–12 / leg"),
        ("Glutes", "Single-leg glute bridge · 3 × 10–15 / side"),
        ("Calves", "Single-leg calf raise off a step · 4 × 8–15 / side · 3 s down, 1 s pause at the bottom · hold 10 kg when easy"),
        ("Core", "Dead bug 3 × 8 / side · side plank 3 × 30–45 s · hollow hold 3 × 20–40 s · reverse crunch 3 × 10–15"),
    ]),
    ("PROGRESSIVE OVERLOAD · tick 'Progression achieved'", YELLOW, [
        ("Ask", "Did I improve something today?"),
        ("Counts", "+1 rep · +1 set · better form · slower eccentric · longer hold · harder variation · less rest · more range · heavier dumbbell"),
        ("Why", "This matters more than 'workout completed'. It's how the 10 kg dumbbell stays useful for months."),
    ]),
    ("20-MINUTE MINIMUM WORKOUT · when a session is missed", ORANGE, [
        ("Rule", "Don't move the whole week around. Do this once, then carry on with the plan."),
        ("4 rounds", "10 push-ups (any version) · 12 goblet squats · 10 rows / side · 30 s plank · 15 calf raises"),
        ("Counts as", "Training / movement. Tick Strength too if it ran 15+ minutes."),
    ]),
    ("CARDIO & STEPS", GREEN, [
        ("Cardio", "2–3 × 20–30 min: brisk or incline walking, cycling, easy jog, low-impact"),
        ("Steps", "Week 1: 7,000 · week 2: 8,000 · week 3+: 9,000 → a 9–10k average"),
        ("Rest day", "At least one lower-intensity day: walking, mobility, light cycling, family activity"),
    ]),
    ("NUTRITION TARGETS", ORANGE, [
        ("Calories", "Start at 1,800–2,000 kcal. Adjust on the 2–3 week trend, not on one day. Never auto-drop to 1,500."),
        ("Protein", "150 g / day · breakfast 30–40 · lunch 35–45 · snack 20–30 · dinner 40–50"),
        ("Whey isolate", "One serving ≈ 25.5 g protein · 113 kcal. Use it to fill gaps."),
        ("Fat", "50–65 g / day. Portion awareness, not elimination."),
        ("Carbs", "The rest of the calories. No carb phobia."),
        ("Water", "2.5–3 L = 500 ml × 5–6. More in heat or after sweaty sessions."),
        ("Caffeine", "If you drink coffee or tea: nothing after 14:00."),
        ("Daily", "Protein in every main meal · 2+ veg · 1–2 fruit · mostly minimally processed · no liquid calories · no unplanned binges"),
    ]),
    ("MAURITIUS-FRIENDLY PLATE · no seafood, no pork", ORANGE, [
        ("Protein", "Chicken · eggs · lean beef · dholl / lentils · beans · yoghurt · cheese · milk · whey"),
        ("Carbs", "Rice · potatoes · oats · bread · roti / farata (watch the portion) · fruit"),
        ("Veg", "Brèdes · cabbage · pumpkin · chou chou · carrots · tomatoes · salad"),
        ("Example day", "Oats + whey + banana (~35 g) · chicken curry, rice, salad (~40 g) · yoghurt or 2 eggs (~20 g) · grilled chicken or beef, potatoes, veg (~45 g)"),
        ("Meal prep", "Sunday 30–60 min: cook a protein, a carb base, chop veg, boil eggs, portion snacks. Make the healthy option the lazy option."),
    ]),
    ("EVENING SHUTDOWN · 20:00 → 22:00", VIOLET, [
        ("20:00", "Work shutdown: what was completed? what remains? tomorrow's first priority? Then stop."),
        ("21:00", "Prep clothes and training kit · review tomorrow's calendar · hygiene · teeth · short prayer · gratitude"),
        ("Screens", "No doomscrolling in bed. Last 30 minutes screen-free."),
        ("21:30", "Read / quiet time"),
        ("22:00", "Lights out"),
    ]),
    ("WORK", ORANGE, [
        ("Before work", "Write the Top 3 priorities."),
        ("During", "2–3 deep-work blocks · move 1–2 min every 45–60 min · 20-20-20 for the eyes"),
        ("Boundary", "Work ends at ~20:00 so the evening belongs to people, not the inbox."),
    ]),
    ("FAITH", YELLOW, [
        ("Daily", "Morning prayer · Scripture · gratitude · evening prayer"),
        ("Weekly", "Church / worship · longer Scripture study · prayer for family · prayer for direction and wisdom"),
        ("Monthly", "Am I becoming more disciplined, loving, truthful, patient and useful?"),
        ("Note", "The tracker measures practice. It doesn't pretend to measure faith."),
    ]),
    ("RELATIONSHIPS & FAMILY", FUCHSIA, [
        ("Fiancée", "10–20 min a day, no phone. 'How are you really doing?' Not every chat needs to be a summit."),
        ("Couple time", "Weekly: dinner, walk, coffee, movie, day trip, private time, a shared activity"),
        ("Son", "Weekly: conversation · shared activity · guidance · positive reinforcement. Connection, not policing."),
        ("Family", "Weekly check-in and one meaningful conversation"),
    ]),
    ("MIND", VIOLET, [
        ("Reading", "20 min a day. Track minutes, not pages. Business · psychology · philosophy · faith · health · strategy · fiction sometimes"),
        ("Learning", "20–30 min of deliberate learning: AI · marketing · business · design · web · sales · leadership"),
        ("Journal (opt.)", "AM: What matters most today? What could derail me? How will I respond?  PM: What did I accomplish? Where did I fall short? What will I do differently?"),
        ("Check-in", "Mood, energy, stress 1–10. Optional prompt: What am I avoiding?"),
        ("Meditation", "5–10 min: breath, contemplative prayer, silence or guided. Consistency before length."),
    ]),
    ("EYES, JAW & DESK · keratoconus-aware", TEAL, [
        ("Eyes", "20-20-20: every ~20 min look ~20 ft away for ~20 s · don't rub your eyes · wear your prescribed correction · notice and ease squinting"),
        ("Check-ups", "Keep regular appointments with your eye specialist. No eye exercises to 'fix' keratoconus."),
        ("Jaw", "Relax it, don't force it into position. If the alignment concern persists, book a dental / orthodontic assessment."),
        ("Desk (weekly)", "Screen near eye level · feet supported · shoulders relaxed · elbows comfortable · alternate sitting and standing"),
    ]),
    ("HOME & LIFE ADMIN", LIME, [
        ("Daily", "Make the bed · 5-minute tidy · dishes and kitchen reset"),
        ("Weekly", "Laundry · room · bathroom · grocery plan · meal prep · financial review · calendar review"),
    ]),
    ("MINIMUM VIABLE DAY · mode M", YELLOW, [
        ("When", "Work explodes, you're tired, travelling, or life throws a chair through the window."),
        ("The list", "Drink water · 10-minute walk · protein as close as possible · 5-minute mobility · prayer · 10 minutes with your fiancée · lights out on time"),
        ("Scoring", "Type M in the day-mode row. The day is scored on these 7 items only. That's it. No guilt."),
    ]),
    ("SICK-DAY RULE · mode S", PINK, [
        ("Don't", "Hard training · aggressive calorie deficit · forced cardio"),
        ("Do", "Rest · hydration · nutrition · sleep · medical care when needed"),
        ("Scoring", "Type S. The day is paused: no score, and it doesn't drag the rolling averages down. Resume training gradually."),
    ]),
    ("WEIGHING, MEASURING, PACE", PINK, [
        ("Weigh-in", "Daily, after the bathroom, before food or drink, similar clothing. Don't react to one number."),
        ("Decide on", "The 7-day average only."),
        ("Every 2 weeks", "Waist · chest · arms · thigh · calf (neck optional)"),
        ("Monthly", "Progress photos: same light, distance, pose and time · front, side, back"),
        ("Pace", "0.4–0.8 kg a week. Flat for 14 days → trim 150–200 kcal or add steps. Above 1%/week repeatedly → eat a bit more."),
        ("Protect muscle", "Watch for strength falling, protein under target, short sleep, very low calories, too much cardio, constant fatigue."),
    ]),
    ("DELIBERATELY NOT INCLUDED", DIM, [
        ("No", "Excessive skincare · jawline or facial exercises · eye exercises to 'fix' keratoconus · excessive stretching · daily calorie-burn targets · punishment workouts · starvation targets · more than one weight metric a day · dozens of tiny habits · streaks that reward unhealthy behaviour"),
        ("Why", "The system should make life simpler, not turn you into the middle manager of your own existence."),
    ]),
]

START_HERE = [
    ("THE 60-SECOND DAILY CHECK-IN", YELLOW, [
        ("1", "Open this month's tab (Oct 2026, Nov 2026…). Today's date is highlighted in yellow."),
        ("2", "Tick the habits you did. Level 3 is bonus only."),
        ("3", "Type the numbers you have in BODY & MIND: weight, steps, protein, water, sleep, energy, stress, mood. Blank is fine."),
        ("4", "Hard day? Type M in the Day mode row. Sick? Type S. Planned rest? R."),
        ("5", "Glance at the Dashboard: Foundation score, rolling 7-day, Smart Coach."),
    ]),
    ("THE TABS", TEAL, [
        ("Dashboard", "Today · this week · this month · coach · system check · pillars · charts. Change the yellow date to look back."),
        ("Week Plan", "Top 3 and tasks for each day, daily donuts and check-ins (task-tracker layout)."),
        ("Month tabs", "The habit grid, scores, body & mind numbers, charts. This is where you log."),
        ("Body", "Profile, goal estimate, measurements every 2 weeks, monthly photos, trend charts."),
        ("Weekly Review", "Sunday, 15–30 min. Numbers fill themselves; write wins, problems and ONE change."),
        ("Monthly Review", "Once a month: body, health, mind, spirit, relationships, business, stop/start/continue."),
        ("Habit Library", "Every habit with level, pillar, target, what counts, and 7/30-day scores."),
        ("Plan & Routines", "The playbook: routines, training, nutrition, MVD, sick days, eye and jaw notes."),
        ("Settings", "Targets and thresholds. Change them here and everything follows."),
        ("Calc · Data", "The engine. Don't type there."),
    ]),
    ("THREE LEVELS", BLUE, [
        ("Level 1", "Non-negotiables. The small daily foundation: morning reset, prayer, mobility, movement, Top 3, fiancée, read/learn, shutdown, lights out (+ protein, water, steps, sleep from your numbers)."),
        ("Level 2", "High value, scored against weekly targets: strength 4×, progression, core 3×, calves 3×, cardio 2×, deep work, breaks, home reset, son, couple time, family, church, meal prep."),
        ("Level 3", "Optional. Never counts against you."),
    ]),
    ("HOW SCORING WORKS", VIOLET, [
        ("Foundation", "Each day gets one score from 9 things: sleep, movement, protein, water, steps, prayer, Top 3 (work days only), fiancée, read/learn."),
        ("Rolling", "The headline numbers are rolling 7-day and 30-day averages, not streaks. One missed day can't wipe anything out."),
        ("Bands", "Green ≥ 85% · Good 70–84% · Needs attention 50–69% · Reset < 50%"),
        ("Modes", "M days are scored on the Minimum Viable Day list. S days are paused."),
        ("Missed vs failed", "The System Check separates 'one off day', 'missed a few habits' and 'the system is too complicated'. It never says 'failed'. When adherence drops, it tells you to simplify, not to try harder."),
    ]),
    ("PRINCIPLES", GREEN, [
        ("1", "Consistency > intensity"),
        ("2", "Progress > perfection"),
        ("3", "Systems > motivation"),
        ("4", "Health > extreme results"),
        ("Note", "This is a personal planning tool, not medical advice. For eyes, jaw or anything persistent, see a professional."),
    ]),
    ("GOOGLE SHEETS SETUP · once", ORANGE, [
        ("1", "Google Drive → New → File upload → Lucas-Life-OS.xlsx → open it → File → Save as Google Sheets."),
        ("2", "In the Google Sheet: Extensions → Apps Script. Paste the contents of life_os_setup.gs, save, click Run on finishSetup and allow access."),
        ("3", "That converts the grids to real checkboxes, applies Montserrat, dark-styles the charts and adds a 'Life OS' menu (Go to today · New week)."),
        ("No script?", "Everything still works: type x in a cell instead of ticking. Or select a grid and use Insert → Checkbox."),
    ]),
]


def build_text(ws, sections, title, subtitle, tab):
    page(ws)
    ws.sheet_properties.tabColor = tab
    widths(ws, {"A": 2, "B": 16, "C": 58, "D": 3, "E": 16, "F": 58, "G": 2})
    put(ws, "B2:F2", title, size=24, bold=True)
    put(ws, "B3:F3", subtitle, size=8, color=MUTED)
    ws.row_dimensions[2].height = 34
    r = 5
    max_row = 5
    for i in range(0, len(sections), 2):
        pair = sections[i:i + 2]
        lens = [len(s[2]) for s in pair]
        rows_needed = max(lens) + 1
        for j, (stitle, col, lines) in enumerate(pair):
            lc, tc = ("B", "C") if j == 0 else ("E", "F")
            put(ws, f"{lc}{r}:{tc}{r}", stitle, size=8, bold=True, color=col, bg=mix(col, PANEL, .16), indent=1)
            for t in range(rows_needed - 1):
                rr = r + 1 + t
                if t < len(lines):
                    lab, txt = lines[t]
                    put(ws, f"{lc}{rr}", lab, size=7, bold=True, color=col, bg=PANEL, v="top", indent=1, wrap=True)
                    put(ws, f"{tc}{rr}", txt, size=8, color=TEXT, bg=PANEL, v="top", wrap=True, indent=1)
                else:
                    paint(ws, f"{lc}{rr}:{tc}{rr}", PANEL)
        for t in range(rows_needed - 1):
            rr = r + 1 + t
            longest = 0
            for stitle, col, lines in pair:
                if t < len(lines):
                    longest = max(longest, len(lines[t][1]))
            ws.row_dimensions[rr].height = max(16, 12 * -(-longest // 66) + 5)
        ws.row_dimensions[r].height = 20
        r += rows_needed + 1
        max_row = r
    paint_bg_outside(ws, f"A1:G{max_row + 2}")


def paint_bg_outside(ws, rng):
    for c in cells(ws, rng):
        if c.fill is None or c.fill.fill_type is None:
            c.fill = PF(BG)


# ======================================================================== APPS SCRIPT
def write_script(path):
    import json as _json
    sheets_order = ["Start Here", "Dashboard", "Week Plan"] + [m[3] for m in MONTHS] + [
        "Body", "Weekly Review", "Monthly Review", "Habit Library", "Plan & Routines", "Settings", "Calc", "Data"]
    js = f"""/**
 * Life OS · finish setup for Google Sheets
 * Extensions → Apps Script → paste this file → Save → run finishSetup once (allow access).
 * Safe to run again at any time: it keeps your ticks and only restyles.
 */
const LIFE_OS = {{
  bg: '#{BG}', panel: '#{PANEL}', panel2: '#{PANEL2}', line: '#{LINE}', text: '#{TEXT}', muted: '#{MUTED}',
  font: '{FONT}',
  checkboxes: {_json.dumps(CHECKBOXES, indent=2, ensure_ascii=False)},
  charts: {_json.dumps(CHART_STYLES, indent=2)},
  monthSheets: {_json.dumps([m[3] for m in MONTHS])},
  order: {_json.dumps(sheets_order)},
}};

function onOpen() {{
  SpreadsheetApp.getUi().createMenu('Life OS')
    .addItem('Go to today', 'goToToday')
    .addItem('New week: clear Week Plan tasks', 'clearWeekPlan')
    .addSeparator()
    .addItem('Finish setup (checkboxes, font, charts)', 'finishSetup')
    .addToUi();
}}

function finishSetup() {{
  const ss = SpreadsheetApp.getActive();
  setupCheckboxes_(ss);
  ss.getSheets().forEach(function (sh) {{
    try {{ sh.getRange(1, 1, sh.getMaxRows(), sh.getMaxColumns()).setFontFamily(LIFE_OS.font); }} catch (e) {{}}
    try {{ sh.setHiddenGridlines(true); }} catch (e) {{}}
    paintOutside_(sh);
  }});
  styleCharts_(ss);
  ['Calc', 'Data'].forEach(function (n) {{ const s = ss.getSheetByName(n); if (s) s.hideSheet(); }});
  SpreadsheetApp.flush();
  ss.toast('Life OS is ready. Use the Life OS menu → Go to today.', 'Life OS', 8);
}}

function setupCheckboxes_(ss) {{
  Object.keys(LIFE_OS.checkboxes).forEach(function (name) {{
    const sh = ss.getSheetByName(name);
    if (!sh) return;
    LIFE_OS.checkboxes[name].forEach(function (a1) {{
      const r = sh.getRange(a1);
      const keep = r.getValues().map(function (row) {{
        return row.map(function (v) {{ return v === true || String(v).toLowerCase() === 'x'; }});
      }});
      r.clearDataValidations();
      r.insertCheckboxes();
      r.setValues(keep);
    }});
  }});
}}

function paintOutside_(sh) {{
  const lastRow = Math.max(sh.getLastRow(), 1), lastCol = Math.max(sh.getLastColumn(), 1);
  const maxRow = sh.getMaxRows(), maxCol = sh.getMaxColumns();
  if (maxRow > lastRow) sh.getRange(lastRow + 1, 1, maxRow - lastRow, maxCol).setBackground(LIFE_OS.bg);
  if (maxCol > lastCol) sh.getRange(1, lastCol + 1, lastRow, maxCol - lastCol).setBackground(LIFE_OS.bg);
}}

function styleCharts_(ss) {{
  const axis = {{ textStyle: {{ color: LIFE_OS.muted, fontName: LIFE_OS.font, fontSize: 10 }},
                 gridlines: {{ color: LIFE_OS.line }}, baselineColor: LIFE_OS.line }};
  Object.keys(LIFE_OS.charts).forEach(function (name) {{
    const sh = ss.getSheetByName(name);
    if (!sh) return;
    const specs = LIFE_OS.charts[name].slice();
    sh.getCharts().forEach(function (ch, i) {{
      const spec = takeSpec_(specs, ch);
      if (!spec) return;
      try {{
        const b = ch.modify()
          .setOption('backgroundColor', LIFE_OS.panel)
          .setOption('chartArea', {{ backgroundColor: LIFE_OS.panel }})
          .setOption('fontName', LIFE_OS.font)
          .setOption('title', '');
        if (spec.colors.length) b.setOption('colors', spec.colors);
        if (spec.kind === 'donut') {{
          b.setOption('pieHole', 0.72).setOption('legend', {{ position: 'none' }})
           .setOption('pieSliceText', 'none').setOption('pieSliceBorderColor', LIFE_OS.panel);
        }} else {{
          const v = JSON.parse(JSON.stringify(axis));
          if (spec.vmin !== undefined) v.viewWindow = {{ min: spec.vmin, max: spec.vmax }};
          if (spec.fmt === 'percent') v.format = 'percent';
          b.setOption('vAxis', v)
           .setOption('hAxis', {{ textStyle: axis.textStyle, gridlines: {{ color: 'transparent' }}, baselineColor: LIFE_OS.line }})
           .setOption('legend', {{ position: 'bottom', textStyle: {{ color: LIFE_OS.muted, fontName: LIFE_OS.font }} }});
          if (spec.kind === 'area') b.setOption('areaOpacity', 0.35).setOption('lineWidth', 2).setOption('legend', {{ position: 'none' }});
          if (spec.kind === 'line' || spec.kind === 'weight') b.setOption('curveType', 'function').setOption('lineWidth', 2);
          if (spec.kind === 'bar') b.setOption('legend', {{ position: 'none' }});
        }}
        sh.updateChart(b.build());
      }} catch (e) {{
        console.log('Chart style skipped on ' + name + ' #' + i + ': ' + e);
      }}
    }});
  }});
}}

// Pick the style whose anchor is closest to where this chart sits.
function takeSpec_(specs, ch) {{
  if (!specs.length) return null;
  let row = 0, col = 0;
  try {{ const ci = ch.getContainerInfo(); row = ci.getAnchorRow(); col = ci.getAnchorColumn(); }} catch (e) {{}}
  let best = 0, bestD = Infinity;
  specs.forEach(function (s, i) {{
    const d = Math.abs(s.row - row) + Math.abs(s.col - col);
    if (d < bestD) {{ bestD = d; best = i; }}
  }});
  return specs.splice(best, 1)[0];
}}

function goToToday() {{
  const ss = SpreadsheetApp.getActive();
  const tz = ss.getSpreadsheetTimeZone();
  const now = new Date();
  const name = Utilities.formatDate(now, tz, 'MMM yyyy');
  const sh = ss.getSheetByName(name) || ss.getSheetByName('Dashboard');
  ss.setActiveSheet(sh);
  if (LIFE_OS.monthSheets.indexOf(name) >= 0) {{
    const day = Number(Utilities.formatDate(now, tz, 'd'));
    sh.setActiveRange(sh.getRange(10, 4 + day));
  }}
}}

function clearWeekPlan() {{
  const ss = SpreadsheetApp.getActive();
  const sh = ss.getSheetByName('Week Plan');
  if (!sh) return;
  const ui = SpreadsheetApp.getUi();
  if (ui.alert('Clear this week\\'s Top 3 and tasks?', ui.ButtonSet.YES_NO) !== ui.Button.YES) return;
  for (let d = 0; d < 7; d++) {{
    const c = 2 + d * 3;
    sh.getRange(31, c, 3, 1).clearContent();
    sh.getRange(35, c, 8, 1).clearContent();
    sh.getRange(31, c + 1, 3, 1).uncheck();
    sh.getRange(35, c + 1, 8, 1).uncheck();
  }}
}}
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(js)


# ======================================================================== DEMO DATA
def fill_demo(wb, today):
    rnd = random.Random(7)
    for k, y, m, sn in MONTHS:
        ws = wb[sn]
        for j in range(1, 32):
            try:
                d = dt.date(y, m, j)
            except ValueError:
                continue
            if d < TRACK_START or d > today:
                continue
            col = dcol(j)
            wd = d.isoweekday()
            sick = rnd.random() < 0.03
            mvd = not sick and rnd.random() < 0.08
            if sick:
                ws[f"{col}{R_MODE}"] = "S"
            elif mvd:
                ws[f"{col}{R_MODE}"] = "M"
            base = 0.55 if mvd else 0.82
            for key, nm, lvl, *_ in HABITS:
                p = {"L1": base, "L2": 0.5, "L3": 0.2}[lvl]
                if key == "str":
                    p = 0.9 if wd in (1, 2, 4, 5) else 0.05
                if key == "prog":
                    continue
                if key in ("core", "calf"):
                    p = 0.8 if wd in (2, 5) else 0.15
                if key in ("church",):
                    p = 0.9 if wd == 7 else 0
                if key in ("prep",):
                    p = 0.8 if wd == 7 else 0
                if key in ("son", "couple", "family"):
                    p = 0.25 if wd >= 5 else 0.05
                if key == "top3" and wd > 5:
                    p = 0.1
                if not sick and rnd.random() < p:
                    ws[f"{col}{HROW[key]}"] = "x"
            if ws[f"{col}{HROW['str']}"].value == "x" and rnd.random() < 0.6:
                ws[f"{col}{HROW['prog']}"] = "x"
            n = (d - TRACK_START).days
            ws[f"{col}{MROW['weight']}"] = round(76.2 - 0.075 * n + rnd.uniform(-0.5, 0.5), 1)
            ws[f"{col}{MROW['steps']}"] = int(rnd.gauss(8800, 1700))
            ws[f"{col}{MROW['protein']}"] = int(rnd.gauss(140, 18))
            ws[f"{col}{MROW['water']}"] = round(rnd.uniform(1.9, 3.1), 1)
            ws[f"{col}{MROW['kcal']}"] = int(rnd.gauss(1920, 140))
            ws[f"{col}{MROW['sleeph']}"] = round(rnd.uniform(6.0, 8.2), 1)
            ws[f"{col}{MROW['energy']}"] = rnd.randint(5, 9)
            ws[f"{col}{MROW['stress']}"] = rnd.randint(2, 7)
            ws[f"{col}{MROW['mood']}"] = rnd.randint(5, 9)
    body = wb["Body"]
    for i in range(27):
        d = TRACK_START + dt.timedelta(days=14 * i)
        if d > today:
            break
        body[f"D{17 + i}"] = round(92 - 0.8 * i + rnd.uniform(-0.3, 0.3), 1)
        body[f"E{17 + i}"] = 101.0
        body[f"H{17 + i}"] = 36.5 + 0.1 * i
        if i % 2 == 0:
            body[f"J{17 + i}"] = "x"
    week = wb["Week Plan"]
    tasks = ["Client proposal", "Pipeline review", "Invoice clients", "Draft homepage copy", "Call accountant",
             "Gym bag ready", "Book eye check-up", "Plan date night", "Groceries"]
    for d in range(7):
        c = CL(2 + d * 3)
        cc = CL(3 + d * 3)
        for r in range(31, 34):
            week[f"{c}{r}"] = rnd.choice(tasks)
            if rnd.random() < 0.6:
                week[f"{cc}{r}"] = "x"
        for r in range(35, 35 + rnd.randint(2, 5)):
            week[f"{c}{r}"] = rnd.choice(tasks)
            if rnd.random() < 0.5:
                week[f"{cc}{r}"] = "x"


# ======================================================================== MAIN
def main():
    global TODAY, TRACK_START
    ap = argparse.ArgumentParser()
    ap.add_argument("--demo", action="store_true", help="fill sample data and fix 'today' (preview only)")
    ap.add_argument("--out", default=os.path.join(HERE, "Lucas-Life-OS.xlsx"))
    args = ap.parse_args()
    demo_today = dt.date(2026, 11, 26)
    if args.demo:
        TODAY = f"DATE({demo_today.year},{demo_today.month},{demo_today.day})"
        TRACK_START = dt.date(2026, 10, 1)
        for i, row in enumerate(SETTINGS):
            if row[1] == "TrackStart":
                SETTINGS[i] = (row[0], row[1], TRACK_START, row[3], row[4])

    wb = Workbook()
    order = ["Start Here", "Dashboard", "Week Plan"] + [m[3] for m in MONTHS] + [
        "Body", "Weekly Review", "Monthly Review", "Habit Library", "Plan & Routines", "Settings", "Calc", "Data"]
    wb.active.title = order[0]
    for nm in order[1:]:
        wb.create_sheet(nm)

    build_settings(wb, wb["Settings"])
    build_data(wb, wb["Data"])
    for k, y, m, sn in MONTHS:
        build_month(wb, wb[sn], k, y, m)
    build_library(wb, wb["Habit Library"])
    build_calc(wb, wb["Calc"])
    build_dashboard(wb, wb["Dashboard"])
    build_week(wb, wb["Week Plan"])
    build_body(wb, wb["Body"])
    build_weekly(wb, wb["Weekly Review"])
    build_monthly(wb, wb["Monthly Review"])
    build_text(wb["Plan & Routines"], ROUTINES, "PLAN & ROUTINES",
               "The playbook behind the ticks. Routines are grouped so each one is a single checkbox.", VIOLET)
    build_text(wb["Start Here"], START_HERE, "LIFE OS · START HERE",
               "Lucas's personal operating system. Consistency > Intensity · Progress > Perfection · "
               "Systems > Motivation · Health > Extreme results", YELLOW)
    wb.active = wb.sheetnames.index("Dashboard")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Dashboard"

    if args.demo:
        fill_demo(wb, demo_today)
    wb.save(args.out)
    if not args.demo:
        write_script(os.path.join(HERE, "life_os_setup.gs"))
    print("saved", args.out)


if __name__ == "__main__":
    main()
